"""Finish the copy on existing template slides: Executive summary, Situation, 1 · What changed, 2 · Network,
Appendix A1 (data cleaning & assumptions) and the closing "Thank you" slide.

Run from the repo root:  uv run --no-project --with python-pptx python slides/build_copyfix.py
Edits the deck (slides/deckkit.py: DECK, or NB_DECK) in place and is safe to rerun: text boxes are rewritten
whole (keeping each box's template paragraph/run formatting), shapes this script adds are named "cf_*" and
removed before they are re-added, and picture crops/sizes are set to absolute values. Slides are found by
kicker, never by index. The title slide's "[Team name]" is left for the team. Numbers are recomputed from
analysis/airline_analysis.sqlite, analysis/tables/alternatives_matrix.csv and analysis/preparation_audit.json;
the asserts stop the build if the data moves away from the wording.
"""
import copy
import csv
import json
import sqlite3

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from deckkit import (BODY, DECK, GREEN, HEAD, INK, LINE, MINUS, MUTED, NAVY, ORANGE, PANEL, ROOT, WHITE, box,
                     find_slide, money, text)

TAG = "cf_"  # prefix of every shape this script adds


# ---------- data
db = sqlite3.connect(ROOT / "analysis/airline_analysis.sqlite")
q = lambda sql, *a: db.execute(sql, a).fetchone()[0]
REV = {p: q("select sum(amount_cents) from transactions where period=?", p) / 1e8 for p in (1, 2)}
COST = {p: q("select sum(amount_cents) from costs where period=?", p) / 1e8 for p in (1, 2)}
REFUND2 = -q("select sum(amount_cents) from transactions where period=2 and type='Refund'") / 1e8
FLIGHTS = {p: q("select count(*) from cost_events where period=?", p) for p in (1, 2)}
MONTH_FL = dict(db.execute("select month, count(*) from cost_events group by month"))
C1, C2 = REV[1] - COST[1], REV[2] - COST[2]
GATE = {y: q("select sum(a) from (select distinct city, gate, year, amount_cents a from gates where year=?)", y) / 1e8
        for y in ("2024-2025", "2025-2026")}
GATE_ROWS = q("select sum(amount_cents) from gates where year='2025-2026'") / 1e8
# 2026-27 leases are missing: repeat each gate's own 2024-25 -> 2025-26 % increase (same method as the gate tables)
GATE_EST = sum(b * b / a for a, b in db.execute(
    "select x.a, y.a from (select distinct city, gate, amount_cents a from gates where year='2024-2025') x "
    "join (select distinct city, gate, amount_cents a from gates where year='2025-2026') y using(city, gate)")) / 1e8
AFTER = C2 - GATE["2025-2026"]
# like-for-like cost per flight: same direction and calendar month in both years, FY25-26 flight weights
LFL = q("with rm as (select route, substr(month,6,2) mm, period, count(*) n, sum(amount_cents)/1e8 c "
        "from cost_events group by route, mm, period) select sum((b.c/b.n - a.c/a.n) * b.n) from rm a join rm b "
        "on a.route=b.route and a.mm=b.mm and a.period=1 and b.period=2")
COST_UP = COST[2] - COST[1]

contrib = {1: {}, 2: {}}
for p in (1, 2):
    for route, cents in db.execute(
            "select route, sum(c) from (select route, amount_cents c from transactions where period=? union all "
            "select route, -amount_cents from costs where period=?) group by route", (p, p)):
        contrib[p][route] = cents / 1e8
P2 = contrib[2]
HUBS = {"YVR", "YYC", "YYZ", "YUL"}
FIVE = {"YUL-YVR", "YVR-YUL", "YHZ-YYZ", "YYZ-YHZ", "YUL-YYZ", "YYZ-YUL", "YOW-YYZ", "YYZ-YOW", "YHZ-YUL", "YUL-YHZ"}
NEW = {"YOW-YUL", "YUL-YOW", "YXU-YYZ", "YYZ-YXU", "YQM-YHZ", "YHZ-YQM"}
pos = [v for v in P2.values() if v > 0]
neg = [v for v in P2.values() if v <= 0]
mtl_van = P2["YUL-YVR"] + P2["YVR-YUL"]
five = sum(P2[r] for r in FIVE)
trunk4 = five - mtl_van  # the 4 doubled trunk markets, 8 directions
new3 = sum(P2[r] for r in NEW)
rest38 = sum(v for r, v in P2.items() if r not in FIVE | NEW)
# monthly margin = (revenue by purchase month - cost by flight month) / revenue
MARGIN = {m: (r - c) / r for m, r, c in db.execute(
    "select t.month, t.v, c.v from (select month, sum(amount_cents) v from transactions group by month) t "
    "join (select month, sum(amount_cents) v from costs group by month) c using(month)")}
decline = {r: P2[r] - contrib[1][r] for r in P2}
top8 = sorted(decline.values())[:8]
TOP8_SHARE = sum(top8) / sum(decline.values())
cls = lambda r: ("hub–hub", "hub–regional", "regional–regional")[2 - sum(a in HUBS for a in r.split("-"))]
rest = {}
for r, v in P2.items():
    if r not in FIVE | NEW:
        rest[cls(r)] = rest.get(cls(r), 0) + v

alt = list(csv.DictReader(open(ROOT / "analysis/tables/alternatives_matrix.csv")))
RS = {r["component"]: r for r in alt if r["alternative"] == "Right-Size & Recover"}
comp = lambda key: float(next(v for k, v in RS.items() if k.startswith(key))["central_musd"])
TRUNK, MTL, HALVE, LEVERS = comp("Remove 2025"), comp("Exit Montreal"), comp("Halve"), comp("Revenue recovery")
TOT = RS["TOTAL"]
# After gates on the deck's basis (REC_SPEC.md): the 2025-26 lease total charged to FY25-26 (AFTER = -58.25). The
# csv's own after_gates_* columns may use another gate basis (FY pro-rated $10.02M gives -58.40), so rebuild them
# from the option's gains: after = AFTER + gates released + gain (low / central / high).
RS_LO, RS_MID, RS_HI = (AFTER + float(TOT["gates_released_musd"]) + float(TOT[k])
                        for k in ("low_musd", "central_musd", "high_musd"))
RS_FLIGHTS = int(TOT["flights_removed"])
audit = json.load(open(ROOT / "analysis/preparation_audit.json"))

# the wording below depends on these; stop rather than publish a sentence the data no longer supports
assert abs(C2 - -48.38) < 0.01 and abs(C1 - -12.36) < 0.01, (C1, C2)
assert abs(AFTER - -58.25) < 0.01 and abs(C2 + REFUND2 - -40.54) < 0.01, (AFTER, C2 + REFUND2)
assert abs(LFL - 15.25) < 0.01 and abs(COST_UP - 45.39) < 0.01, (LFL, COST_UP)
assert (MONTH_FL["2025-02"], MONTH_FL["2025-03"]) == (792, 1157)
assert len(P2) == 54 and len(pos) == 26 and abs(five - -48.70) < 0.01 and five < C2
assert 0.30 < mtl_van / sum(neg) < 0.36 and abs(TOP8_SHARE - 0.713) < 0.002
assert abs(TRUNK + MTL + HALVE - 54.57) < 0.01 and abs(LEVERS - 5.2) < 0.01, (TRUNK, MTL, HALVE, LEVERS)
assert (round(RS_LO, 1), round(RS_MID, 2), round(RS_HI, 1)) == (-16.7, 1.52, 5.9), (RS_LO, RS_MID, RS_HI)
# the top 8 decliners are exactly both directions of the 4 named markets (Halifax–Montréal is 9th/10th)
assert set(sorted(decline, key=decline.get)[:8]) == FIVE - {"YHZ-YUL", "YUL-YHZ"}
assert len(FIVE | NEW) == 16 and abs(trunk4 - -28.99) < 0.01 and abs(new3 - -10.25) < 0.01
assert abs(rest38 - 10.57) < 0.01 and abs(five + new3 - -58.95) < 0.01
assert abs(MARGIN["2025-02"] - 0.045) < 0.001 and MARGIN["2025-03"] < 0
assert all(v < 0 for m, v in MARGIN.items() if "2025-03" <= m <= "2026-08"), "margin turned positive again"
# "doubled": each of the five markets flew at least 1.85x (1.89-2.08x) as many flights Sep–Feb FY25-26 as Sep–Feb FY24-25
SEPFEB = dict(((r, p), n) for r, p, n in db.execute(
    "select route, period, count(*) from cost_events where substr(month,6,2) in ('09','10','11','12','01','02') "
    "group by route, period"))
assert all(SEPFEB[r, 2] >= 1.85 * SEPFEB[r, 1] for r in FIVE), {r: (SEPFEB[r, 1], SEPFEB[r, 2]) for r in FIVE}
FL_UP = FLIGHTS[2] / FLIGHTS[1] - 1
REV_UP = REV[2] / REV[1] - 1
SCHED = TRUNK + MTL + HALVE


# ---------- helpers
def usd(x, sign=False):
    """$ millions, rounded half-up like the rest of the deck: usd(-58.25) -> '−$58.3M', usd(5.2, True) -> '+$5.2M'."""
    v = money(x, sign=sign)
    return f"{v[0]}${v[1:]}M" if v[0] in (MINUS, "+") else f"${v}M"


# template text styles (explicit, so a rerun never inherits formatting from text this script wrote)
TITLE = {"font": HEAD, "size": 24, "bold": True, "color": NAVY}
SOURCE = {"font": BODY, "size": 9, "bold": False, "color": MUTED}
GREY = RGBColor.from_string("9CA3AF")


def fill(shape, paras, base, align=None, space_after=None):
    """Rewrite a template text box in place: paragraphs keep the box's first paragraph properties, runs keep its
    first run's properties, then get `base` (font, size, bold, color) and per-run overrides. `paras` items are a
    str or a list of (text, overrides); a first-run override "space_after" (pt) sets that paragraph's spacing.
    Default spacing after: 0 for titles and source lines, 6pt otherwise."""
    if space_after is None:
        space_after = 0 if base is TITLE or base is SOURCE else 6
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    proto_p = copy.deepcopy(p0._p)
    for r in proto_p.findall(qn("a:r")) + proto_p.findall(qn("a:br")) + proto_p.findall(qn("a:fld")):
        proto_p.remove(r)
    r0 = p0.runs[0]._r if p0.runs else None
    proto_rpr = r0.find(qn("a:rPr")) if r0 is not None else None
    proto_rpr = copy.deepcopy(proto_rpr) if proto_rpr is not None else None
    body = tf._txBody
    for p in list(tf.paragraphs):
        body.remove(p._p)
    for para in paras:
        body.append(copy.deepcopy(proto_p))
        p = tf.paragraphs[-1]
        if align is not None:
            p.alignment = align
        runs = [(para, {})] if isinstance(para, str) else para
        p.space_after = Pt(runs[0][1].get("space_after", space_after))
        for t, o in runs:
            r = p.add_run()
            if proto_rpr is not None:
                old = r._r.find(qn("a:rPr"))
                if old is not None:
                    r._r.remove(old)
                r._r.insert(0, copy.deepcopy(proto_rpr))
            r.text = t
            st = {**base, **o}
            f = r.font
            f.name, f.size, f.bold, f.italic = st["font"], Pt(st["size"]), st["bold"], False
            f.color.rgb = st["color"]


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def drop_added(slide, media=False):
    """Remove the shapes this script added; with media=True also the template picture. Pictures and charts drop
    their relationship too, so reruns do not pile up orphaned chart/image parts."""
    for sh in list(slide.shapes):
        is_pic, is_chart = sh.shape_type == 13, getattr(sh, "has_chart", False) and sh.has_chart
        if not (sh.name.startswith(TAG) or (media and is_pic)):
            continue
        rid = sh._element.blipFill.blip.rEmbed if is_pic else sh._element.chart_rId if is_chart else None
        sh._element.getparent().remove(sh._element)
        if rid:
            slide.part.drop_rel(rid)


def drop_stripe(slide):
    """The template's thin orange edge stripe on the SO WHAT panel (Rectangle 9, 0.08in wide); house style has none."""
    for sh in list(slide.shapes):
        if sh.name == "Rectangle 9" and sh.width < Inches(0.1):
            sh._element.getparent().remove(sh._element)


def hide_label(point):
    """Delete one point's data label (python-pptx has no API for a per-point <c:delete>)."""
    dl = point.data_label._get_or_add_dLbl()
    for child in list(dl):
        if child.tag != qn("c:idx"):
            dl.remove(child)
    etree.SubElement(dl, qn("c:delete")).set("val", "1")


def chart_layout(chart, x, y, w, h):
    """Fix the inner plot area (fractions of the chart frame) so bar positions are known; transparent background."""
    cs = chart._chartSpace
    pa = cs.find(qn("c:chart")).find(qn("c:plotArea"))
    old = pa.find(qn("c:layout"))
    if old is not None:
        pa.remove(old)
    lay = etree.Element(qn("c:layout"))
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"), ("x", x), ("y", y), ("w", w),
                     ("h", h)):
        etree.SubElement(ml, qn("c:" + tag)).set("val", str(val))
    pa.insert(0, lay)
    sp = cs.find(qn("c:spPr"))
    if sp is None:
        sp = etree.Element(qn("c:spPr"))
        cs.find(qn("c:chart")).addnext(sp)
    for child in list(sp):
        sp.remove(child)
    etree.SubElement(sp, qn("a:noFill"))
    etree.SubElement(etree.SubElement(sp, qn("a:ln")), qn("a:noFill"))


def quiet_axes(ch, vmin, vmax, reverse=False):
    ca = ch.category_axis
    ca.reverse_order = reverse
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.color.rgb = LINE
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(10.5)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minimum_scale, va.maximum_scale = vmin, vmax


def circle(slide, x, y, d, num, name):
    c = box(slide, x, y, d, d, fill=ORANGE, shape=MSO_SHAPE.OVAL, name=name)
    tf = c.text_frame
    tf.text = num
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = HEAD, Pt(12), True, WHITE
    return c


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


B = {"bold": True, "color": NAVY}  # bold lead-in run


# ---------- executive summary
def build_exec(s):
    drop_added(s)
    fill(shape(s, "TextBox 3"),
         [f"NorthByte flew {FL_UP:.0%} more flights for {REV_UP:.0%} more revenue; Right-Size & Recover reaches "
          "break-even"], TITLE)
    fill(shape(s, "TextBox 5"),
         ["Source: NorthByte transactions, costs & gates, Sep 2024–Aug 2026; team analysis. After gates = 2025-26 "
          f"leases (${GATE['2025-2026']:.2f}M). Impact: central case, FY25-26 basis."], SOURCE)
    fill(shape(s, "TextBox 8"), [
        f"• FY25-26 revenue +{REV_UP:.1%} to {usd(REV[2])}; flight operating cost +{COST[2] / COST[1] - 1:.1%} "
        f"to {usd(COST[2])}",
        f"• Flight contribution fell from {usd(C1)} to {usd(C2)} before gate leases ({usd(AFTER)} after). "
        f"The case's ~$40M matches our loss before refunds ({usd(C2 + REFUND2)})",
        f"• Monthly margin turned negative in Mar 2025 (Feb still +{MARGIN['2025-02']:.1%}), when the second "
        f"trunk rotations began: flights {MONTH_FL['2025-02']:,} → {MONTH_FL['2025-03']:,} a month",
        f"• ⅔ of the {usd(COST_UP, True)} cost growth is more flights and route mix ({usd(COST_UP - LFL)}); "
        "only ⅓ is higher cost per flight",
        f"• The 16 directions added or doubled in 2025 lose ~${-(five + new3):.0f}M (4 trunk markets "
        f"{usd(trunk4)}, Montréal–Vancouver {usd(mtl_van)}, 3 new markets {usd(new3)}); the other 38 earn "
        f"{usd(rest38, True)}",
        "• Refunds, Rewards and ancillaries are real, but an order of magnitude smaller than the "
        f"{usd(C2 - C1)} decline",
    ], {"font": BODY, "size": 14, "bold": False, "color": INK})

    # right column: the placeholder box becomes a plain panel
    panel = shape(s, "Rectangle 9")
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    ln = panel._element.spPr.find(qn("a:ln"))
    if ln is not None and ln.find(qn("a:prstDash")) is not None:
        ln.remove(ln.find(qn("a:prstDash")))
    panel.line.fill.background()
    px, pw = panel.left / 914400, panel.width / 914400  # 6.85, 5.88
    head = shape(s, "TextBox 10")
    head.left, head.top, head.width, head.height = Inches(px + 0.15), Inches(1.85), Inches(pw - 0.3), Inches(0.62)
    head.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    fill(head, [[("WHAT WE RECOMMEND", {"font": HEAD, "size": 11, "bold": True, "color": ORANGE, "space_after": 2})],
                [("Right-Size & Recover  ", {"font": HEAD, "size": 14, "bold": True, "color": NAVY}),
                 ("highest-scoring of four options", {"size": 11, "color": MUTED})]],
         {"font": BODY, "size": 12, "bold": False, "color": INK}, align=PP_ALIGN.LEFT)

    recs = [
        ("1", "Right-size the schedule", usd(SCHED, True),
         f"Remove the 2025 added flights on 4 trunk markets ({usd(TRUNK, True)}); exit Montréal–Vancouver "
         f"({usd(MTL, True)}); halve the 3 new markets, with a kill trigger ({usd(HALVE, True)})"),
        ("2", "Recover revenue", usd(LEVERS, True),
         "Restore trunk fares, rebook-first for St. John's disruptions, shift OTA bookings to direct, "
         "promote hub connections"),
        ("3", "Run a route-level P&L", "fixes root cause",
         "Reviewed monthly, with add/remove triggers: flights return only on proven demand, peak days first"),
    ]
    x, w = px + 0.15, pw - 0.3
    y = 2.55
    for i, (num, title, stat, body) in enumerate(recs):
        h = 1.02 if i == 0 else 0.84
        circle(s, x, y + 0.02, 0.34, num, f"{TAG}rec{num}_num")
        text(s, x + 0.47, y, 3.2, 0.3, [title], size=13, bold=True, color=NAVY, font=HEAD, name=f"{TAG}rec{num}_head")
        text(s, x + 3.6, y, w - 3.6, 0.3, [stat], size=13 if stat[0] == "+" else 10, bold=stat[0] == "+",
             color=GREEN if stat[0] == "+" else MUTED, font=HEAD if stat[0] == "+" else BODY, align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE, name=f"{TAG}rec{num}_stat")
        text(s, x + 0.47, y + 0.32, w - 0.47, h - 0.36, [body], size=10.5, name=f"{TAG}rec{num}_body")
        y += h + 0.06

    cy = 5.48
    box(s, x, cy, w, 1.15, fill=WHITE, line=LINE, name=f"{TAG}impact_card")
    text(s, x + 0.15, cy + 0.08, w - 0.3, 0.26, [[("EXPECTED IMPACT  ", {"font": HEAD, "size": 10, "bold": True,
                                                                        "color": ORANGE}),
                                                  ("contribution after gates, FY25-26 basis", {"size": 10,
                                                                                               "color": MUTED})]],
         name=f"{TAG}impact_label")
    text(s, x + 0.15, cy + 0.34, w - 0.3, 0.36, [[(usd(AFTER), {"color": NAVY}),
                                                 ("  →  ", {"color": MUTED}),
                                                 (f"≈ {usd(RS_MID, True)}", {"color": GREEN})]],
         size=18, bold=True, font=HEAD, name=f"{TAG}impact_value")
    text(s, x + 0.15, cy + 0.72, w - 0.3, 0.4,
         [f"Range {usd(RS_LO, True)} to {usd(RS_HI, True)}. All 20 airports kept; "
          f"{RS_FLIGHTS / FLIGHTS[2]:.0%} fewer flights. Target: break-even after gates in FY27-28."],
         size=10.5, name=f"{TAG}impact_note")

    notes(s, f"""EXECUTIVE SUMMARY. ~45 sec. Answer first.
Say: "NorthByte flew {FL_UP:.0%} more flights for {REV_UP:.0%} more revenue, and its flight loss nearly quadrupled. The loss sits in the flights added in 2025. We recommend Right-Size & Recover: cut those flights, recover revenue, and run the airline route by route. That takes NorthByte from −$58.3M to about break-even after gates, with all 20 airports kept."
Left column (what we found), all FY25-26 vs FY24-25 from analysis/airline_analysis.sqlite:
- Revenue ${REV[1]:.2f}M → ${REV[2]:.2f}M; flight costs ${COST[1]:.2f}M → ${COST[2]:.2f}M. Flights {FLIGHTS[1]:,} → {FLIGHTS[2]:,} (+{FL_UP:.1%}).
- Contribution {C1:.2f} → {C2:.2f} before gates; {AFTER:.2f} after the ${GATE['2025-2026']:.2f}M of 2025-26 gate leases (each physical gate once, split by departing flights per the case rule). The case's "about $40M" matches our loss with refunds (${REFUND2:.2f}M) left out: {C2 + REFUND2:.2f}. Its "four times the year before" matches only with refunds included ({C2 / C1:.1f}×). If asked: the case's two figures fit different definitions; do not call either an error.
- Monthly margin (revenue by purchase month, costs by flight month): +{MARGIN['2025-02']:.1%} in Feb 2025, {MARGIN['2025-03']:.1%} in Mar, negative every month to Aug 2026. The case says losses ran from February; on our data February is still positive. Second rotations: Montréal–Toronto and Ottawa–Toronto from Mar 2025; Montréal–Vancouver, Halifax–Toronto and Halifax–Montréal from Apr; Montréal–Ottawa launched 1 Mar. Flights Feb → Mar 2025 {MONTH_FL['2025-02']} → {MONTH_FL['2025-03']:,} (+{MONTH_FL['2025-03'] / MONTH_FL['2025-02'] - 1:.0%} as monthly counts; about +32% per day, since February is short). Full year +{FL_UP:.1%}.
- Cost growth split: like-for-like (same route, same calendar month, FY25-26 flight weights) ${LFL:.2f}M = ⅓ (+7.1% per flight); volume and route mix ${COST_UP - LFL:.2f}M = ⅔ (HANDOFF.md; slide 8).
- 16 directions = the 4 doubled trunk markets (Montréal–Toronto, Halifax–Toronto, Halifax–Montréal, Ottawa–Toronto; 8 directions) {trunk4:.2f}, Montréal–Vancouver {mtl_van:.2f}, the 3 new markets (London–Toronto, Montréal–Ottawa, Moncton–Halifax) {new3:.2f}: total {five + new3:.2f}; the other 38 directions +{rest38:.2f} (slides/OUTLINE.md "Verification notes").
- Not-the-problem items against the −$36.0M decline: refunds −$1.15M, Points + Cash at most +$1.33M, ancillaries +$0.29M (OUTLINE.md).
Right column (what we recommend) from analysis/tables/alternatives_matrix.csv, Right-Size & Recover rows; it must match the Recommendation slide:
- Schedule ${SCHED:.1f}M = {TRUNK:.1f} (2,265 added trunk flights; +$14.5M if no revenue moves to the kept flights) + {MTL:.1f} (Montréal–Vancouver, 620 flights, 0 connecting passengers) + {HALVE:.1f} (halve London–Toronto, Montréal–Ottawa, Moncton–Halifax, 1,411 flights; still about −$2.6M after halving, so pair it with the 0.8 revenue/cost kill trigger by Q4 FY26-27).
- Revenue levers ${LEVERS:.1f}M central ($2.2–8.3M): trunk yield, St. John's rebook-first, OTA to direct, hub connections, community fares, all-in fare bundles.
- After gates: {AFTER:.1f} → +{RS_MID:.2f} central (range {RS_LO:.1f} to +{RS_HI:.1f}: 0% to 100% of revenue on removed flights moving to kept flights). Schedule changes alone reach about {usd(AFTER + SCHED)}. It is break-even on recorded costs, not a margin: aircraft ownership and overhead are not in the data. {RS_FLIGHTS:,} of {FLIGHTS[2]:,} flights removed ({RS_FLIGHTS / FLIGHTS[2]:.0%}).""")


# ---------- situation
def build_situation(s):
    card = {"font": BODY, "size": 12, "bold": False, "color": MUTED}
    fill(shape(s, "TextBox 12"), [f"FY25-26 flight contribution before gates ({usd(AFTER)} after)"], card)
    fill(shape(s, "TextBox 18"), ["New markets launched Feb–Apr 2025: London, Moncton, Montréal–Ottawa"], card)
    fill(shape(s, "TextBox 5"), ["Source: FirstByte 2026 case package; NorthByte transactions, costs & gates "
                                 "(Sep 2024–Aug 2026), team analysis"], SOURCE)
    notes(s, f"""SITUATION / KEY QUESTION. ~40 sec.
Set up the case from the brief: domestic carrier, 20 airports, 4 hubs (YVR, YYC, YYZ, YUL), fiscal year Sep–Aug. In early 2025 NorthByte opened 3 new markets and added a second rotation on five trunk markets. Finance reviewed costs only at company level, by category, never route by route (case p.10); that gap is our angle. End on the key question. The values (safety, fairness to customers, connecting communities) are the criteria we use to choose between alternatives later.
Numbers: revenue ${REV[2]:.2f}M (+{REV_UP:.1%}). −$48M = {C2:.2f} before gates (revenue minus fuel, crew, maintenance and food, analysis/airline_analysis.sqlite); {AFTER:.2f} after ${GATE['2025-2026']:.2f}M of 2025-26 gate leases. The case's "about $40M" matches our loss with refunds excluded ({C2 + REFUND2:.1f}). 20 airports and 54 directional routes counted in the cost file. New markets' first flights: London–Toronto 1 Feb, Montréal–Ottawa 1 Mar, Moncton–Halifax 1 Apr 2025 (the case says "spring 2025").""")


# ---------- 1 · what changed
CITY = {"YUL": "Montréal", "YVR": "Vancouver", "YYZ": "Toronto", "YYC": "Calgary", "YHZ": "Halifax", "YOW": "Ottawa"}


def build_changed(s):
    drop_added(s, media=True)  # template PNG (matplotlib, "Period 1/2") is replaced by a native waterfall
    drop_stripe(s)
    cat = {k: (q("select sum(amount_cents) from costs where period=2 and type=?", k)
               - q("select sum(amount_cents) from costs where period=1 and type=?", k)) / 1e8
           for k in ("Fuel", "Crew", "Maintenance", "Food")}
    assert all(v > 0 for v in cat.values()) and abs(sum(cat.values()) - COST_UP) < 0.01, cat

    text(s, 0.6, 1.9, 8.0, 0.3, [[("Contribution bridge, FY24-25 → FY25-26  ", {"bold": True, "color": NAVY}),
                                  ("$M, before gate leases", {"color": MUTED, "size": 11})]],
         size=13, name=f"{TAG}chart_title")
    # waterfall as stacked columns: an invisible base, then one coloured series per bar type
    steps = [("FY24-25", "total", C1), ("Revenue", "gain", REV[2] - REV[1])] + \
            [(k, "cost", -v) for k, v in cat.items()] + [("FY25-26", "total", C2)]
    run, base, bars = 0.0, [], {"total": [], "gain": [], "cost": []}
    for name, kind, v in steps:
        if kind == "total":
            top, length, run = 0.0, v, v
        else:
            top, length, run = max(run, run + v), -abs(v), run + v
        assert top <= 0 and top + length <= 0, (name, top, length)  # the whole bridge sits below zero
        base.append(round(top, 4) if top else None)
        for k in bars:
            bars[k].append(round(length, 4) if k == kind else None)
    assert abs(run - C2) < 1e-6, run
    cd = CategoryChartData()
    cd.categories = [n for n, _, _ in steps]
    cd.add_series("base", base)
    series = [("total", NAVY, '0.0;"−"0.0'), ("gain", GREEN, '"+"0.0;"+"0.0'), ("cost", ORANGE, '0.0;"−"0.0')]
    for k, _, _ in series:
        cd.add_series(k, bars[k])
    fx, fy, fw, fh = 0.6, 2.3, 8.0, 4.5
    lx, ly, lw, lh = 0.02, 0.02, 0.96, 0.86
    vmin = -55
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(fx), Inches(fy), Inches(fw), Inches(fh), cd)
    gf.name = f"{TAG}bridge"
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.name, ch.font.size, ch.font.color.rgb = BODY, Pt(10), INK
    chart_layout(ch, lx, ly, lw, lh)
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 45, 100
    b = plot.series[0]
    b.format.fill.background()
    b.format.line.fill.background()
    for ser, (_, color, fmt) in zip(list(plot.series)[1:], series):
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = color
        dls = ser.data_labels
        dls.show_value = True
        dls.number_format, dls.number_format_is_linked = fmt, False
        dls.position = XL_LABEL_POSITION.CENTER
        dls.font.name, dls.font.size, dls.font.bold, dls.font.color.rgb = BODY, Pt(11), True, WHITE
    food = len(steps) - 2
    hide_label(plot.series[3].points[food])  # too short for an inside label: written below the bar instead
    quiet_axes(ch, vmin, 0)
    xc = fx + fw * (lx + lw * (food + 0.5) / len(steps))
    yb = fy + fh * (ly + lh * (-C2) / -vmin)  # the food bar ends at the FY25-26 total
    text(s, xc - 0.4, yb + 0.03, 0.8, 0.26, [money(-cat["Food"])], size=11, bold=True, align=PP_ALIGN.CENTER,
         name=f"{TAG}food_label")

    fill(shape(s, "TextBox 11"), [
        f"• Revenue {usd(REV[2] - REV[1], True)} (+{REV_UP:.1%})",
        f"• Flight costs {usd(COST_UP, True)} (+{COST[2] / COST[1] - 1:.1%})",
        f"• Every cost line grew, mostly from {FL_UP:.0%} more flights: fuel {usd(cat['Fuel'], True)}, crew "
        f"{usd(cat['Crew'], True)}, maintenance {usd(cat['Maintenance'], True)}",
        f"• Gate leases (${GATE['2025-2026']:.2f}M) push the loss to {usd(AFTER)}",
        f"• The case's “about $40M” matches this loss before refunds ({usd(C2 + REFUND2)})",
    ], {"font": BODY, "size": 13, "bold": False, "color": INK})
    fill(shape(s, "TextBox 5"), ["Source: NorthByte transactions & costs, FY24-25 vs FY25-26; contribution = revenue "
                                 "less fuel, crew, maintenance, food; gate leases 2025-26 (each gate once)"],
         SOURCE)
    lfl_cat = {k: q("with f as (select route, substr(month,6,2) mm, period, count(*) n from cost_events group by 1,2,3), "
                    "k as (select route, substr(month,6,2) mm, period, sum(amount_cents)/1e8 c from costs where type=? "
                    "group by 1,2,3), rm as (select k.*, f.n from k join f using(route, mm, period)) "
                    "select sum((b.c/b.n - a.c/a.n)*b.n) from rm a join rm b on a.route=b.route and a.mm=b.mm "
                    "and a.period=1 and b.period=2", k) for k in ("Fuel", "Crew", "Maintenance")}
    assert all(lfl_cat[k] < cat[k] / 2 for k in lfl_cat), lfl_cat  # "mostly from more flights" holds per line
    ratio = (COST[2] / COST[1] - 1) / REV_UP
    notes(s, f"""1 · WHAT CHANGED (guiding Q1). ~50 sec.
Say: "Revenue grew ${REV[2] - REV[1]:.1f}M, or {REV_UP:.1%}. Flight costs grew ${COST_UP:.1f}M, or {COST[2] / COST[1] - 1:.1%}: {ratio:.1f} times as fast. A ${-C1:.1f}M loss became ${-C2:.1f}M before gates, and ${-AFTER:.1f}M once gate leases are charged." Then hand over to why costs rose: about ⅔ is more flights and route mix, ⅓ higher cost per flight.
Bridge (analysis/airline_analysis.sqlite, FY24-25 → FY25-26, before gates): {C1:.2f} + revenue {REV[2] - REV[1]:+.2f} − fuel {cat['Fuel']:.2f} − crew {cat['Crew']:.2f} − maintenance {cat['Maintenance']:.2f} − food {cat['Food']:.2f} = {C2:.2f}. "{ratio:.1f}× faster" compares growth rates (+{COST[2] / COST[1] - 1:.1%} vs +{REV_UP:.1%}); in dollars costs grew {COST_UP / (REV[2] - REV[1]):.1f}× as much as revenue.
"Mostly from more flights": flights {FLIGHTS[1]:,} → {FLIGHTS[2]:,} (+{FL_UP:.1%}); like-for-like cost per flight (same route and calendar month) explains only fuel ${lfl_cat['Fuel']:.1f}M, crew ${lfl_cat['Crew']:.1f}M and maintenance ${lfl_cat['Maintenance']:.1f}M of these increases; the rest is volume and route mix (slide 8).
Gates: ${GATE['2025-2026']:.2f}M for the 2025-26 lease year with each physical gate counted once (the raw rows sum to ${GATE_ROWS:.1f}M because shared gates repeat on every route); after gates {AFTER:.2f}, split across departing flights per the case rule.
Case reconciliation: the case's "about $40M" matches our loss without refunds ({C2 + REFUND2:.2f}; refunds ${REFUND2:.2f}M). With refunds included the loss is {C2 / C1:.1f}× the prior year, which matches the case's "four times".""")


# ---------- 2 · network
def build_network(s):
    drop_added(s, media=True)  # template PNG showed airport codes only; replaced by a native chart with city names
    drop_stripe(s)
    fill(shape(s, "TextBox 3"),
         [f"Five markets doubled in 2025 lose {usd(-five)}, as much as the whole network; {len(pos)} of {len(P2)} "
          "routes profit"], TITLE)

    text(s, 0.6, 1.9, 8.0, 0.3, [[("Contribution by route, FY25-26  ", {"bold": True, "color": NAVY}),
                                  ("$M before gates: the 10 worst and 6 best of 54 routes", {"color": MUTED, "size": 11})]],
         size=13, name=f"{TAG}chart_title")
    ranked = sorted(P2, key=P2.get)
    shown = ranked[:10] + ranked[-6:][::-1]
    assert set(ranked[:10]) == FIVE, ranked[:10]  # the 10 worst directions are exactly the five doubled markets
    best = ranked[-6:]
    assert all(cls(r) == "hub–hub" for r in best), best
    name = lambda r: f"{CITY[r[:3]]} → {CITY[r[4:]]}"
    cd = CategoryChartData()
    cd.categories = [name(r) for r in shown]
    cd.add_series("Contribution", [round(P2[r], 4) for r in shown])
    fx, fy, fw, fh = 0.6, 2.3, 8.0, 4.55
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(fx), Inches(fy), Inches(fw), Inches(fh), cd)
    gf.name = f"{TAG}routes"
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.name, ch.font.size, ch.font.color.rgb = BODY, Pt(10), INK
    chart_layout(ch, 0.25, 0.01, 0.71, 0.98)
    plot = ch.plots[0]
    plot.gap_width = 40
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '"+"0.0;"−"0.0', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.name, dl.font.size, dl.font.bold, dl.font.color.rgb = BODY, Pt(10.5), True, INK
    ser = plot.series[0]
    ser.invert_if_negative = False
    for i, r in enumerate(shown):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = ORANGE if r in FIVE else GREEN
    quiet_axes(ch, -12.5, 3, reverse=True)

    fill(shape(s, "TextBox 11"), [
        "• The five: Montréal–Vancouver, Halifax–Toronto, Montréal–Toronto, Ottawa–Toronto, Halifax–Montréal. "
        f"The first four (both ways) = {TOP8_SHARE:.0%} of the {usd(C2 - C1)} decline",
        f"• Montréal–Vancouver alone: {usd(mtl_van)}, ~⅓ of all route losses",
        f"• Profitable routes ({usd(sum(pos), True)}) can't offset losers ({usd(sum(neg))})",
        f"• Outside the 2025 expansion every route type is near break-even or better: hub–hub "
        f"{usd(rest['hub–hub'], True)}, hub–regional {usd(rest['hub–regional'], True)}, regional–regional "
        f"{usd(rest['regional–regional'])}",
    ], {"font": BODY, "size": 13, "bold": False, "color": INK})
    fill(shape(s, "TextBox 5"), ["Source: NorthByte transactions & costs, FY25-26 (decline vs FY24-25); contribution "
                                 "before gates; routes are directional. Hubs: YVR, YYC, YYZ, YUL (case)."], SOURCE)
    n_rest = {k: sum(1 for r in P2 if r not in FIVE | NEW and cls(r) == k) for k in rest}
    notes(s, f"""2 · NETWORK (guiding Q2). ~60 sec.
Say: "The five markets that got a second rotation in 2025 (Montréal–Vancouver, Halifax–Toronto, Montréal–Toronto, Ottawa–Toronto, Halifax–Montréal) lose ${-five:.1f}M, as much as the whole network. Montréal–Vancouver alone is a third of all route losses and carries no connecting passengers. {len(pos)} routes still make money, so the rest of the network is carrying the expansion."
Numbers (analysis/airline_analysis.sqlite, FY25-26, before gates, 54 directional routes; all deduplicated transactions):
- The five markets = 10 directions, {five:.2f}; network {C2:.2f}. They are the ten orange bars, the 10 worst of 54; each flew about twice as many flights Sep–Feb as a year earlier. The six green bars are the best routes: Toronto, Calgary and Vancouver hub pairs.
- Montréal–Vancouver {mtl_van:.2f} = {mtl_van / sum(neg):.0%} of the {sum(neg):.2f} lost by the {len(neg)} negative directions.
- {len(pos)} positive directions earn +{sum(pos):.2f}. (Dropping the 759 out-of-window transactions gives the same 26 routes: +$12.37M / −$61.04M.)
- Top 8 decliners (both directions of Montréal–Vancouver, Halifax–Toronto, Montréal–Toronto, Ottawa–Toronto): {sum(top8):.2f} of the {C2 - C1:.2f} decline = {TOP8_SHARE:.1%}. Halifax–Montréal is 9th and 10th.
- Hub split (hubs YVR, YYC, YYZ, YUL), outside the 16 expansion directions: hub–hub {n_rest['hub–hub']} directions +{rest['hub–hub']:.2f} (all profitable); hub–regional {n_rest['hub–regional']} directions +{rest['hub–regional']:.2f}; regional–regional {n_rest['regional–regional']} directions {rest['regional–regional']:.2f} (Halifax–St. John's, Edmonton–Yellowknife). Including the expansion, every class loses (hub–hub −21.2, hub–regional −23.9, regional–regional −3.3), so what separates winners from losers is whether a market was expanded, not hub vs regional.
- Both directions share the same sign in all 27 markets: 13 of 27 markets are profitable.
- After gates (case rule, 2025-26 leases): 24 of 54 directions positive; the five markets −51.5 of {AFTER:.1f}. Issue sizes stay before gates because gate leases do not fall when flights are cut.""")


# ---------- appendix A1
def build_a1(s):
    a = audit
    fill(shape(s, "TextBox 7"), [
        [("Route codes: ", B), ("'YYZYHZ', 'yvr-yeg' etc. normalized to XXX-XXX in all three files. Each direction "
                                "is its own route: 54 directions, 27 markets, 20 airports", {})],
        [("Duplicates: ", B), (f"{a['removed_duplicate_rows']:,} repeated transaction IDs removed (all fields "
                               f"identical), {MINUS}${a['removed_duplicate_revenue_cents'] / 1e8:.2f}M; "
                               f"{a['deduplicated_transaction_rows']:,} transactions remain", {})],
        [("Dates: ", B), ("ISO, YYYY/MM/DD and DD/MM vs MM/DD formats; ambiguous ones resolved from transaction-ID "
                          f"order; {a['date_fallback_rows']} unresolved (${a['date_fallback_gross_cents'] / 1e5:.1f}k, "
                          "immaterial)", {})],
        [("Categories & tiers: ", B), ("aliases and capitalization standardized (e.g. Website → Direct Web); blanks "
                                       "kept as 'Unknown', not assumed to be Non-Member", {})],
        [("Flights: ", B), (f"one flight = route + timestamp with exactly one Crew, Food, Fuel and Maintenance row "
                            f"({a['cost_rows']:,} rows = {a['cost_rows'] // 4:,} flights). The Cost_Date time is "
                            "read as the departure time: slots cluster at 06:30–08:00 and 12:00–18:10, so 12:00 or "
                            "later = afternoon flight", {})],
        [("Revenue timing: ", B), ("Transaction_Date is the purchase date and Cost_Date the flight date, so revenue "
                                   "is compared by period and season, not matched to flights. One ticket transaction "
                                   "= one passenger on one flight (a load proxy; there are no seat counts)", {})],
        [("Contribution: ", B), ("revenue (tickets, extras and change fees, less refunds) minus fuel, crew, "
                                 "maintenance and food. Aircraft ownership, overhead and taxes are not in the data, "
                                 "so it is not net income", {})],
        [("Gates: ", B), (f"each physical gate counted once (34 gates: ${GATE['2024-2025']:.2f}M → "
                          f"${GATE['2025-2026']:.2f}M, +{GATE['2025-2026'] / GATE['2024-2025'] - 1:.1%}; the raw rows "
                          f"sum to ${GATE_ROWS:.1f}M). After-gates results charge the 2025-26 leases to FY25-26 and "
                          "split each airport's total across all its departing flights (case rule). 2026-27 leases "
                          f"are missing: estimated at ${GATE_EST:.2f}M by repeating each gate's own increase", {})],
        [("Periods & connections: ", B), ("FY24-25 = Sep 2024–Aug 2025, FY25-26 = Sep 2025–Aug 2026; two years "
                                          "cannot prove recurring seasonality. 12,644 bookings connect through a "
                                          "hub; 268 more are return trips", {})],
    ], {"font": BODY, "size": 12, "bold": False, "color": INK})
    fill(shape(s, "TextBox 5"), ["Source: NorthByte transactions, costs & gates (Sep 2024–Aug 2026); team data "
                                 "preparation audit and assumptions log"], SOURCE)
    notes(s, f"""APPENDIX A1: DATA CLEANING & ASSUMPTIONS. Q&A only.
Use for "how did you handle X?" questions.
Sources: analysis/preparation_audit.json ({a['raw_transaction_rows']:,} raw transaction rows → {a['deduplicated_transaction_rows']:,}; {a['removed_duplicate_rows']:,} duplicates worth ${a['removed_duplicate_revenue_cents'] / 1e8:.2f}M; {a['date_fallback_rows']} fallback dates worth ${a['date_fallback_gross_cents'] / 100:,.0f}; {a['cost_rows']:,} cost rows; {a['booking_count']:,} bookings) and analysis/ASSUMPTIONS.md (1 duplicates, 2 dates, 4–5 money and contribution, 6 categories, 8 flights, 10 routes, 16 Cost_Date time = departure time).
Assumption 16: the data dictionary only says "the date and time of the flight". We read it as departure time because each route repeats fixed slots in two groups (06:30–08:00 and 12:00–18:10). Any cut-off from 08:01 to 12:00 gives the same split, and if the time were arrival the two groups would still separate, so the afternoon-flight finding does not depend on it.
Gates: lease totals from the gates table with each (city, gate, year) once: ${GATE['2024-2025']:.4f}M and ${GATE['2025-2026']:.4f}M; row sum ${GATE_ROWS:.4f}M. The headline after-gates loss ({AFTER:.2f}) charges the 2025-26 lease total to FY25-26. The 2026-27 estimate (${GATE_EST:.3f}M) repeats each gate's 2024-25 → 2025-26 percentage increase, the same method and total as the teammate's gate forecast table; using it with fiscal-year pro-rating moves the loss only to about −$58.4M (analysis/key_issues_verification.md).
Connections: 12,912 bookings hold tickets on two routes; 12,644 connect through a hub and 268 are return trips (HANDOFF.md).""")


# ---------- closing
def build_close(s):
    fill(shape(s, "TextBox 3"),
         [f"Right-Size & Recover: {usd(AFTER)} to break-even, all 20 airports kept  ·  Questions?"],
         {"font": BODY, "size": 20, "bold": False, "color": GREY})
    notes(s, f"""CLOSE. ~10 sec.
Say: "Fly the flights that pay, recover the revenue NorthByte is leaving behind, and manage the network route by route. That takes the loss after gates from ${-AFTER:.1f}M to about break-even (+${RS_MID:.1f}M central) without leaving any of the 20 airports." Then open Q&A (5 min); the appendix slides are for answering questions.
Source: analysis/tables/alternatives_matrix.csv, Right-Size & Recover TOTAL gains on the 2025-26 lease basis (after gates central {RS_MID:.2f}, range {RS_LO:.2f} to {RS_HI:.2f}).""")


def main():
    prs = Presentation(DECK)
    build_exec(find_slide(prs, "EXECUTIVE SUMMARY"))
    build_situation(find_slide(prs, "SITUATION"))
    build_changed(find_slide(prs, "1 · WHAT CHANGED"))
    build_network(find_slide(prs, "2 · NETWORK"))
    build_a1(find_slide(prs, "APPENDIX A1 · DATA CLEANING & ASSUMPTIONS"))
    build_close(find_slide(prs, "Thank you"))
    prs.save(DECK)
    print(f"built copyfix slides: contribution {C1:+.2f} -> {C2:+.2f} (after gates {AFTER:+.2f}); "
          f"five markets {five:+.2f}; positive routes {len(pos)}; Right-Size {RS_MID:+.2f} ({RS_LO:+.2f} to {RS_HI:+.2f})")


if __name__ == "__main__":
    main()
