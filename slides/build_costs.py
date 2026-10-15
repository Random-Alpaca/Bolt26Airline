"""Fill the "1 · WHY COSTS ROSE" slide: the +$45.4M flight-cost rise split route by route into
volume & route mix (about two-thirds) and like-for-like cost per flight (about one-third).

Run from the repo root:  uv run --no-project --with python-pptx python slides/build_costs.py
Edits the deck (slides/deckkit.py: DECK, or NB_DECK) in place and is safe to rerun: the slide is reset to its
frame (top bar, kicker, title, footer) before content is added. Every number is recomputed from
analysis/airline_analysis.sqlite (costs table: one row per flight and category) and checked against the
verified figures in slides/OUTLINE.md "Verification notes" / analysis/key_issues_verification.md.

Method: for each route and calendar month flown in both years (612 route-months), like-for-like =
FY25-26 cost − FY25-26 flights × FY24-25 cost per flight on that route and month. Volume & route mix =
total growth − like-for-like (more flights at each route's own FY24-25 rate, plus route-months new in FY25-26).
"""
import sqlite3
from collections import defaultdict
from decimal import ROUND_HALF_UP, Decimal

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from deckkit import (BODY, DECK, HEAD, INK, LINE, MINUS, MUTED, NAVY, ORANGE, PANEL, ROOT, WHITE, box, find_slide,
                     reset, text)

KICKER = "1 · WHY COSTS ROSE"
CATS = ["Fuel", "Crew", "Maintenance", "Food"]
GROUPS = {
    "trunk": {"YHZ-YYZ", "YYZ-YHZ", "YUL-YYZ", "YYZ-YUL", "YOW-YYZ", "YYZ-YOW", "YHZ-YUL", "YUL-YHZ"},
    "mtl": {"YUL-YVR", "YVR-YUL"},
    "new": {"YOW-YUL", "YUL-YOW", "YXU-YYZ", "YYZ-YXU", "YQM-YHZ", "YHZ-YQM"},
}
MONTHS = ["09", "10", "11", "12", "01", "02", "03", "04", "05", "06", "07", "08"]


# ---------- data
def group(route):
    return next((g for g, rs in GROUPS.items() if route in rs), "rest")


db = sqlite3.connect(ROOT / "analysis/airline_analysis.sqlite")
C, N = defaultdict(float), {}  # C[(route, mm, period, category)] dollars; N[(route, mm, period)] flights
for route, mm, period, cat, dollars, n in db.execute(
        "select route, substr(month, 6, 2), period, type, sum(amount_cents) / 100.0, count(*) "
        "from costs group by 1, 2, 3, 4"):
    C[(route, mm, period, cat)] = dollars
    N[(route, mm, period)] = n  # one cost row per category per flight
K1 = {(r, m) for r, m, p in N if p == 1}
K2 = {(r, m) for r, m, p in N if p == 2}
MATCHED = K1 & K2
assert not K1 - K2 and len(MATCHED) == 612, (len(K1 - K2), len(MATCHED))


def growth(cats, keep=lambda r: True):
    return (sum(C[(r, m, 2, c)] for r, m in K2 if keep(r) for c in cats)
            - sum(C[(r, m, 1, c)] for r, m in K1 if keep(r) for c in cats))


def lfl(cats, keep=lambda r: True, months=None):
    """(dollar effect, % change) of cost per flight on matched route-months, at FY25-26 flight weights."""
    keys = [(r, m) for r, m in MATCHED if keep(r) and (months is None or m in months)]
    now = sum(C[(r, m, 2, c)] for r, m in keys for c in cats)
    base = sum(C[(r, m, 1, c)] / N[(r, m, 1)] * N[(r, m, 2)] for r, m in keys for c in cats)
    return now - base, now / base - 1


def per_flight(cat, period, keep=lambda r: True):
    return (sum(v for (r, m, p, c), v in C.items() if p == period and c == cat and keep(r))
            / sum(n for (r, m, p), n in N.items() if p == period and keep(r)))


TOTAL = {p: sum(v for k, v in C.items() if k[2] == p) for p in (1, 2)}
FLIGHTS = {p: sum(n for k, n in N.items() if k[2] == p) for p in (1, 2)}
GROWTH = TOTAL[2] - TOTAL[1]
LFL, LFL_PCT = lfl(CATS)
VOL = GROWTH - LFL
CAT = {c: (growth([c]) - lfl([c])[0], lfl([c])[0], lfl([c])[1]) for c in CATS}  # vol&mix, lfl $, lfl %
GRP = {g: growth(CATS, lambda r, g=g: group(r) == g) - lfl(CATS, lambda r, g=g: group(r) == g)[0]
       for g in ("trunk", "mtl", "new", "rest")}  # volume & mix by route group
GRP_LFL = lfl(CATS, lambda r: group(r) != "rest")[0]  # like-for-like on the 16 added/doubled directions
GRP_FLIGHTS = {g: tuple(sum(n for (r, m, p), n in N.items() if p == q and group(r) == g) for q in (1, 2))
               for g in ("trunk", "mtl", "new", "rest")}
AVG = {c: per_flight(c, 2) / per_flight(c, 1) - 1 for c in ("Fuel", "Crew")}
EXNEW = {c: per_flight(c, 2, lambda r: r not in GROUPS["new"]) / per_flight(c, 1, lambda r: r not in GROUPS["new"]) - 1
         for c in ("Fuel", "Crew")}
EXNEW_LFL = {c: lfl([c], lambda r: r not in GROUPS["new"])[1] for c in ("Fuel", "Crew")}  # +8.8% / +8.6%


def exnew_mix(cat):
    """Mix effect on the per-flight average without the new markets, by route group, in % points of the FY24-25
    average: (FY25-26 flight share − FY24-25 flight share) × (route-month FY24-25 rate − average) / average."""
    keys = [(r, m) for r, m in K1 if r not in GROUPS["new"]]
    n1, n2 = (sum(N[(r, m, p)] for r, m in keys) for p in (1, 2))
    avg = sum(C[(r, m, 1, cat)] for r, m in keys) / n1
    out = defaultdict(float)
    for r, m in keys:
        out[group(r)] += (N[(r, m, 2)] / n2 - N[(r, m, 1)] / n1) * (C[(r, m, 1, cat)] / N[(r, m, 1)] / avg - 1) * 100
    return out, avg


MIX = {c: exnew_mix(c) for c in ("Fuel", "Crew")}
NEW_LFL = {c: lfl([c], lambda r: r in GROUPS["new"])[1] for c in ("Fuel", "Crew")}  # new markets' matched months
MTL_RATIO = per_flight("Fuel", 1, lambda r: r in GROUPS["mtl"]) / MIX["Fuel"][1]  # ~4x the ex-new average fuel
NEW_KEYS = sorted(K2 - K1)  # route-months flown in FY25-26 but not FY24-25
assert len(NEW_KEYS) == 36 and {r for r, _ in NEW_KEYS} == GROUPS["new"]
CPF = {p: TOTAL[p] / FLIGHTS[p] for p in (1, 2)}
NEW_CPF = sum(C[(r, m, 2, c)] for r, m in K2 if r in GROUPS["new"] for c in CATS) / GRP_FLIGHTS["new"][1]
OTHER_CPF = sum(C[(r, m, 2, c)] for r, m in K2 if r not in GROUPS["new"] for c in CATS) / (
    FLIGHTS[2] - GRP_FLIGHTS["new"][1])
YOY = {m: lfl(CATS, months={m})[1] for m in MONTHS}  # like-for-like year-over-year gap by calendar month

# verified figures (slides/OUTLINE.md "Verification notes"; analysis/key_issues_verification.md)
assert round(GROWTH / 1e6, 2) == 45.39 and round(LFL / 1e6, 2) == 15.25 and round(VOL / 1e6, 2) == 30.14
assert round(LFL_PCT * 100, 1) == 7.1 and round(CAT["Fuel"][2] * 100, 1) == 8.7 and round(CAT["Crew"][2] * 100, 1) == 8.5
assert round(AVG["Fuel"] * 100, 2) == 3.09 and round(AVG["Crew"] * 100, 2) == 4.05
assert round(EXNEW["Fuel"] * 100, 2) == 9.24 and round(EXNEW["Crew"] * 100, 2) == 9.60
assert FLIGHTS == {1: 13072, 2: 15729}
assert round(GRP_LFL / 1e6, 1) == 6.5  # the part of the $15.3M already inside slide 17's issue 1 route losses
assert round(EXNEW_LFL["Fuel"] * 100, 1) == 8.8 and round(EXNEW_LFL["Crew"] * 100, 1) == 8.6
assert GRP_FLIGHTS["trunk"] == (3391, 4530) and GRP_FLIGHTS["mtl"] == (447, 620)
TRUNK_ADDED, TRUNK_ADDED_COST, MTL_LOSS = 2265, 28.5, 19.7  # slides/OUTLINE.md "Verification notes" (before gates)


def m1(x):
    """$M to one decimal, half-up."""
    return str(Decimal(str(x / 1e6)).quantize(Decimal("0.1"), ROUND_HALF_UP))


def pct(x, nd=1):
    s = str(Decimal(str(abs(x) * 100)).quantize(Decimal(1).scaleb(-nd), ROUND_HALF_UP))
    return (MINUS if x < 0 else "+") + s + "%"


def pts(x):
    """Signed percentage points to one decimal, half-up, e.g. +2.4 / −2.2."""
    return pct(x / 100)[:-1]


def usd(x, nd=2):
    """Signed $M for the notes, e.g. −$0.17M."""
    return (MINUS if x < 0 else "") + f"${abs(x) / 1e6:.{nd}f}M"


def split_round(parts, total):
    """Round to 0.1 so the parts add to the rounded total (largest remainder): 30.14 + 15.25 -> 30.1 + 15.3 = 45.4."""
    tenths = [x / 1e5 for x in parts]
    out = [int(t) for t in tenths]
    need = int(Decimal(str(total / 1e5)).quantize(Decimal("1"), ROUND_HALF_UP)) - sum(out)
    for i in sorted(range(len(parts)), key=lambda i: tenths[i] - out[i], reverse=True)[:need]:
        out[i] += 1
    return [o / 10 for o in out]


VOL_M, LFL_M = split_round([VOL, LFL], GROWTH)  # 30.1, 15.3 (LFL is $15.25M; rounded so the parts add to $45.4M)
assert (VOL_M, LFL_M) == (30.1, 15.3)
CARD1_SUM = sum(Decimal(m1(GRP[g])) for g in ("trunk", "mtl", "new")) - Decimal(m1(-GRP["rest"]))  # 30.2 as shown
assert abs(CARD1_SUM - Decimal(str(VOL_M))) <= Decimal("0.1"), CARD1_SUM


# ---------- slide
def hide_label(point):
    """Delete one point's data label (python-pptx has no API for a per-point <c:delete>)."""
    dl = point.data_label._get_or_add_dLbl()
    for child in list(dl):
        if child.tag != qn("c:idx"):
            dl.remove(child)
    etree.SubElement(dl, qn("c:delete")).set("val", "1")


def chart_layout(chart, x, y, w, h):
    """Fix the inner plot area (fractions of the chart frame) so bar positions are known, and clear the chart
    background so the panel shows through."""
    cs = chart._chartSpace
    pa = cs.find(qn("c:chart")).find(qn("c:plotArea"))
    old = pa.find(qn("c:layout"))
    if old is not None:
        pa.remove(old)
    lay = etree.Element(qn("c:layout"))
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"), ("x", x), ("y", y), ("w", w),
                     ("h", h)):
        etree.SubElement(ml, qn("c:" + tag)).set("val", str(val))
    pa.insert(0, lay)
    sp = cs.find(qn("c:spPr"))
    if sp is None:
        sp = etree.Element(qn("c:spPr"))
        cs.find(qn("c:chart")).addnext(sp)
    for child in list(sp):
        sp.remove(child)
    etree.SubElement(sp, qn("a:noFill"))
    etree.SubElement(etree.SubElement(sp, qn("a:ln")), qn("a:noFill"))


def drop_media(slide):
    """Remove pictures and charts together with their relationships, so the old chart image and the charts of
    earlier runs do not pile up in the file (reset() only removes the shapes)."""
    for sh in list(slide.shapes):
        if sh.shape_type == 13:  # PICTURE
            rid = sh._element.blipFill.blip.rEmbed
        elif sh.has_chart:
            rid = sh._element.chart_rId
        else:
            continue
        sh._element.getparent().remove(sh._element)
        slide.part.drop_rel(rid)


def build_costs(s):
    drop_media(s)
    reset(s, KICKER,
          "Two-thirds of the $45M cost rise is more flying; higher cost per flight is the other third",
          "Source: NorthByte flight costs, FY24-25 vs FY25-26 (28,801 flights); team analysis. Like-for-like = same "
          "route and calendar month (612 route-months).",
          NOTES)
    text(s, 0.6, 1.62, 12.13, 0.3, [[
        (f"Flight costs rose ${m1(GROWTH)}M ({pct(GROWTH / TOTAL[1])}) to ${m1(TOTAL[2])}M", {"bold": True, "color": NAVY}),
        (". Each route compared with itself, same month a year earlier:", {})]],
        size=14)

    # left: the chart
    px, py, pw, ph = 0.6, 2.08, 6.3, 4.82
    box(s, px, py, pw, ph, fill=PANEL)
    text(s, px + 0.2, py + 0.1, pw - 0.4, 0.5, [
        [("Where the +$45.4M came from, $M", {"bold": True, "color": NAVY})],
        [("FY24-25 → FY25-26. % in brackets = like-for-like change in cost per flight", {"size": 10, "color": MUTED})]],
        size=12)
    series = [("Volume & route mix (more flights)", NAVY), ("Like-for-like cost per flight", ORANGE)]
    lx = px + 0.2
    for name, color in series:  # legend drawn as shapes: fixed position, matches the card colours
        box(s, lx, py + 0.74, 0.15, 0.15, fill=color)
        text(s, lx + 0.22, py + 0.68, 2.75, 0.26, [name], size=10.5)
        lx += 3.0
    cats = [f"All costs ({pct(LFL_PCT)})"] + [f"{c} ({pct(CAT[c][2])})" for c in CATS]
    vol = [VOL_M] + [round(CAT[c][0] / 1e6, 1) for c in CATS]
    like = [LFL_M] + [max(0.0, round(CAT[c][1] / 1e6, 1)) for c in CATS]  # food: -$0.015M is drawn as 0
    assert round(sum(vol[1:]), 1) == VOL_M and round(sum(like[1:]), 1) == LFL_M, (vol, like)
    cd = CategoryChartData()
    cd.categories = cats
    for (name, _), vals in zip(series, (vol, like)):
        cd.add_series(name, vals)
    fx, fy, fw, fh = px + 0.1, py + 0.98, pw - 0.2, 2.95  # chart frame
    lx_, ly_, lw_, lh_ = 0.25, 0.02, 0.71, 0.96  # inner plot area, fractions of the frame
    vmax = 47
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(fx), Inches(fy), Inches(fw), Inches(fh), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.name, ch.font.size, ch.font.color.rgb = BODY, Pt(10), INK
    chart_layout(ch, lx_, ly_, lw_, lh_)
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 40, 100
    for ser, (_, color) in zip(plot.series, series):
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = color
        dls = ser.data_labels  # series level, so the per-point deletes below do not switch the series off
        dls.show_value = True
        dls.number_format, dls.number_format_is_linked = "0.0", False
        dls.position = XL_LABEL_POSITION.CENTER
        dls.font.name, dls.font.size, dls.font.bold, dls.font.color.rgb = BODY, Pt(11), True, WHITE
        hide_label(ser.points[len(cats) - 1])  # food bar is too short for an inside label
    ca = ch.category_axis
    ca.reverse_order = True
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.color.rgb = LINE
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(10.5)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minimum_scale, va.maximum_scale = 0, vmax
    # food total, placed just past the end of its bar (positions follow from the fixed plot layout)
    band = fh * lh_ / len(cats)
    yc = fy + fh * ly_ + band * (len(cats) - 0.5)
    xe = fx + fw * lx_ + fw * lw_ * (vol[-1] + like[-1]) / vmax
    text(s, xe + 0.06, yc - 0.13, 0.6, 0.26, [m1(CAT["Food"][0] + CAT["Food"][1])], size=11, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, px + 0.2, py + 4.02, pw - 0.4, 0.75, [
        [("Volume & route mix: ", {"bold": True, "color": NAVY}),
         ("FY25-26 flights at their own route's FY24-25 cost per flight in the same month (new route-months in full).", {})],
        [("Like-for-like: ", {"bold": True, "color": ORANGE}),
         ("the change in cost per flight on the same route and month.", {})]], size=10, color=MUTED)

    # right: three cards
    cx, cw, chh, gap = 7.2, 5.53, 1.52, 0.13
    n16 = sum(len(v) for v in GROUPS.values())
    cards = [
        ("1", f"${VOL_M}M", NAVY, "more flying: two-thirds",
         f"Flights +20% ({FLIGHTS[1]:,} → {FLIGHTS[2]:,}). All of it is on the {n16} directions added or doubled "
         f"in 2025: 4 doubled trunk markets ${m1(GRP['trunk'])}M, Montréal–Vancouver ${m1(GRP['mtl'])}M, 3 new markets "
         f"${m1(GRP['new'])}M; other 38 {MINUS}${m1(-GRP['rest'])}M.",
         "→ Fix: right-size; recorded cost falls with each flight cut"),
        ("2", f"${LFL_M}M", ORANGE, "costlier flights: one-third",
         f"Same route, same month: {pct(LFL_PCT)} per flight (fuel {pct(CAT['Fuel'][2])}, crew {pct(CAT['Crew'][2])}, "
         f"maintenance {pct(CAT['Maintenance'][2])}, food flat). Unit costs rose on every route in monthly steps "
         "from Jan to Jul 2025, then flattened.",
         "→ Fix: fuel prices and the pilot agreement are given; fly fewer flights"),
        ("3", "3% / 4%", NAVY, "the case's figures hide it",
         f"Network-average fuel and crew per flight rose only {pct(AVG['Fuel'])} / {pct(AVG['Crew'])} (all costs "
         f"{pct(CPF[2] / CPF[1] - 1)}): the 3 new markets add many cheap flights (~${NEW_CPF / 1e3:.1f}k each vs "
         f"~${OTHER_CPF / 1e3:.1f}k elsewhere). Without them: fuel {pct(EXNEW['Fuel'])}, crew {pct(EXNEW['Crew'])}.",
         "→ Fix: track cost per flight route by route (Issue 3)"),
    ]
    y = 2.08
    for num, stat, scol, head, body, fix in cards:
        box(s, cx, y, cw, chh, fill=WHITE, line=LINE)
        circ = box(s, cx + 0.15, y + 0.12, 0.42, 0.42, fill=ORANGE, shape=MSO_SHAPE.OVAL)
        circ.text_frame.text = num
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = 0
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].runs[0]
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = HEAD, Pt(14), True, WHITE
        text(s, cx + 0.72, y + 0.08, cw - 0.87, 0.36, [[(stat, {"size": 18, "color": scol}),
                                                        ("   " + head, {"size": 13, "color": NAVY})]],
             bold=True, font=HEAD)
        text(s, cx + 0.72, y + 0.47, cw - 0.87, 0.72, [body], size=10.5)
        text(s, cx + 0.72, y + chh - 0.31, cw - 0.87, 0.24, [fix], size=10, bold=True, color=ORANGE)
        y += chh + gap


NOTES = f"""1 · WHY COSTS ROSE. ~60 sec.
Say: "Costs rose $45M. We compared every route with itself in the same month a year earlier. Two-thirds, $30M, is simply more flying: 20% more flights, and all of that extra flying is on the 16 directions NorthByte added or doubled in 2025. One-third, $15M, is each flight costing more: fuel up almost 9%, crew up 8.5%. NorthByte can't set fuel prices, and crew rates are fixed by the new pilot agreement. What it decides is how many flights it flies."
1. Volume & route mix ${VOL_M}M (${VOL / 1e6:.2f}M): FY25-26 flights at each route's own FY24-25 cost per flight, same calendar month; route-months with no FY24-25 flights count in full ({len(NEW_KEYS)} route-months: the new markets' Sep–Mar months, which they had not yet flown a year earlier; they launched Feb–Apr 2025). By group: 4 doubled trunk markets (Montréal–Toronto, Halifax–Toronto, Halifax–Montréal, Ottawa–Toronto) ${GRP['trunk'] / 1e6:.2f}M, Montréal–Vancouver ${GRP['mtl'] / 1e6:.2f}M, 3 new markets ${GRP['new'] / 1e6:.2f}M, other 38 directions {usd(GRP['rest'])} ({GRP_FLIGHTS['rest'][0]:,} → {GRP_FLIGHTS['rest'][1]:,} flights). On the card the rounded parts add to ${CARD1_SUM}M, not ${VOL_M}M, only by rounding; unrounded they add to ${VOL / 1e6:.2f}M.
   Why the trunk shows ${m1(GRP['trunk'])}M here but ${TRUNK_ADDED_COST}M on Key issues and the Recommendation: ${m1(GRP['trunk'])}M is the year-over-year increase. FY24-25 already contained the doubled flights from Mar/Apr to Aug 2025 (trunk flights {GRP_FLIGHTS['trunk'][0]:,} → {GRP_FLIGHTS['trunk'][1]:,}). The {TRUNK_ADDED:,} added trunk flights cost ${TRUNK_ADDED_COST}M in FY25-26: the Recommendation's +${TRUNK_ADDED_COST}M central, which assumes their revenue moves to the kept flights. Likewise Montréal–Vancouver's ${m1(GRP['mtl'])}M is extra cost from {GRP_FLIGHTS['mtl'][0]:,} → {GRP_FLIGHTS['mtl'][1]:,} flights; the Recommendation's +${MTL_LOSS}M is the route's whole FY25-26 loss (revenue minus cost, before gates).
2. Like-for-like ${LFL_M}M (${LFL / 1e6:.2f}M, {pct(LFL_PCT, 2)}; shown as 15.3 so the parts add to $45.4M, as on Key issues): fuel ${CAT['Fuel'][1] / 1e6:.2f}M ({pct(CAT['Fuel'][2], 2)}), crew ${CAT['Crew'][1] / 1e6:.2f}M ({pct(CAT['Crew'][2], 2)}), maintenance ${CAT['Maintenance'][1] / 1e6:.2f}M ({pct(CAT['Maintenance'][2], 2)}), food {usd(CAT['Food'][1], 3)} ({pct(CAT['Food'][2], 2)}).
   Timing: the like-for-like gap is {pct(YOY['09'])} in Sep, {pct(YOY['12'])} in Dec, {pct(YOY['03'])} in Mar, {pct(YOY['06'])} in Jun and {pct(YOY['07'])} / {pct(YOY['08'])} in Jul / Aug. Read: unit costs climbed ~13–14% in monthly steps from Jan to Jul 2025 on every route (standard deviation across routes under 2 points in every month), then stayed flat. Only Jul–Aug is directly observed as flat. It started in Jan, before the March expansion, and hit unchanged routes just as much, so it is not a side effect of the expansion. Do not call it a one-time step.
   Basis matters, footnote if asked: $15.3M at FY25-26 route-month weights; $12.3M on identical flights; $26–28M at run-rate vs fall-2024 unit costs.
3. The case's "fuel +3%, crew +4% per flight" are network averages: {pct(AVG['Fuel'], 2)} / {pct(AVG['Crew'], 2)} (all costs {pct(CPF[2] / CPF[1] - 1, 2)}, ${CPF[1]:,.0f} → ${CPF[2]:,.0f} per flight). The 3 new markets' flights cost ~${NEW_CPF:,.0f} each vs ~${OTHER_CPF:,.0f} elsewhere, which pulls the average down. Excluding them: {pct(EXNEW['Fuel'], 2)} / {pct(EXNEW['Crew'], 2)}.
   Why those sit above like-for-like ({pct(CAT['Fuel'][2])} / {pct(CAT['Crew'][2])}): they are plain averages, so they also carry route mix. On the routes left, like-for-like is {pct(EXNEW_LFL['Fuel'])} / {pct(EXNEW_LFL['Crew'])} (the new markets' own matched months rose only {pct(NEW_LFL['Fuel'])} / {pct(NEW_LFL['Crew'])}, which pulls the network figure down). On top, the mix shifts toward Montréal–Vancouver ({GRP_FLIGHTS['mtl'][0]:,} → {GRP_FLIGHTS['mtl'][1]:,} flights at ~{MTL_RATIO:.0f}x the average fuel cost per flight): {pts(MIX['Fuel'][0]['mtl'])} / {pts(MIX['Crew'][0]['mtl'])} points (fuel / crew). More trunk flights, which are cheaper than average, take {pts(MIX['Fuel'][0]['trunk'])} / {pts(MIX['Crew'][0]['trunk'])}. Net mix {pts(sum(MIX['Fuel'][0].values()))} / {pts(sum(MIX['Crew'][0].values()))} points.
   A company-level, by-category view (case p.10) could not see either effect: this is Issue 3.
Link to Key issues (Issue 2, "each flight earns less and costs more"): it carries the same $15.3M. Do NOT add it to Issue 1: about ${GRP_LFL / 1e6:.1f}M of the like-for-like rise sits on the 16 added directions and is already inside their route losses.
If a judge cites "86% volume" (older draft, and the teammate notebook's $39.2M / $6.2M split): that values every extra flight at the FY24-25 network-average cost (${CPF[1]:,.0f}), but many are cheap new-market flights, so it overstates volume and hides the per-flight rise. Route by route it is two-thirds / one-third.
Sources: analysis/airline_analysis.sqlite, costs table (recomputed by slides/build_costs.py); slides/OUTLINE.md "Verification notes"; analysis/key_issues_verification.md. Costs only, before gates (gate leases +2.2%, mostly fixed)."""


def main():
    prs = Presentation(DECK)
    build_costs(find_slide(prs, KICKER))
    prs.save(DECK)
    print(f"built {KICKER}: growth {GROWTH / 1e6:.2f}M = volume & mix {VOL / 1e6:.2f}M + like-for-like "
          f"{LFL / 1e6:.2f}M ({LFL_PCT:+.2%}); 16 directions LFL {GRP_LFL / 1e6:.2f}M")


if __name__ == "__main__":
    main()
