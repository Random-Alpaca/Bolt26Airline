"""Fill the "2 · THE EXPANSION" slide: added trunk frequency vs the 3 new markets vs everything else.

Run from the repo root:  uv run --no-project --with python-pptx python slides/build_expansion.py
Edits the deck (slides/deckkit.py: DECK, or NB_DECK) in place and is safe to rerun: the slide is reset to its
frame (top bar, kicker, title, footer) and its old charts are dropped before content is added. Every number
is recomputed here from analysis/airline_analysis.sqlite (read-only).

Comparison window: Sep–Feb FY24-25 (2024-09..2025-02, one rotation on the trunk) vs Sep–Feb FY25-26
(2025-09..2026-02, two rotations). Revenue = all transaction types net of refunds, by purchase month; cost and
flights from cost_events (one flight = route + timestamp with 4 cost rows), by flight date; before gates.
"""
import sqlite3
from decimal import ROUND_HALF_UP, Decimal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from deckkit import (BODY, DECK, HEAD, INK, LINE, MINUS, MUTED, NAVY, ORANGE, PANEL, ROOT, WHITE, box, find_slide,
                     label, money, reset, text)

KICKER = "2 · THE EXPANSION"
BEFORE = RGBColor.from_string("B8C2D1")  # Sep 24–Feb 25 bars

# ---------- route groups (directional route codes as stored in the database)
MAR = ["YUL-YYZ", "YYZ-YUL", "YOW-YYZ", "YYZ-YOW"]                      # doubled in Mar 2025
APR = ["YHZ-YYZ", "YYZ-YHZ", "YHZ-YUL", "YUL-YHZ", "YUL-YVR", "YVR-YUL"]  # doubled in Apr 2025
TRUNK = MAR + APR
NEW = ["YXU-YYZ", "YYZ-YXU", "YOW-YUL", "YUL-YOW", "YQM-YHZ", "YHZ-YQM"]
A, B = ("2024-09", "2025-02"), ("2025-09", "2026-02")

db = sqlite3.connect(f"file:{ROOT / 'analysis/airline_analysis.sqlite'}?mode=ro", uri=True)
ALL = [r for (r,) in db.execute("select distinct route from costs order by route")]
OTHER = [r for r in ALL if r not in TRUNK + NEW]
assert len(ALL) == 54 and len(OTHER) == 38
assert db.execute("select count(*) from cost_events where category_count != 4").fetchone()[0] == 0


def agg(routes, win):
    q = ",".join(f"'{r}'" for r in routes)
    rev, tix, trev = db.execute(
        f"select sum(amount_cents), sum(type = 'Ticket'), sum(case when type = 'Ticket' then amount_cents end) "
        f"from transactions where route in ({q}) and month between ? and ?", win).fetchone()
    fl, cost = db.execute(f"select count(*), sum(amount_cents) from cost_events where route in ({q}) "
                          f"and month between ? and ?", win).fetchone()
    return {"rev": rev / 1e8, "tix": tix, "fare": trev / tix / 100, "fl": fl, "cost": cost / 1e8,
            "contrib": (rev - cost) / 1e8, "tpf": tix / fl}


G = {k: (agg(v, A), agg(v, B)) for k, v in (("trunk", TRUNK), ("new", NEW), ("other", OTHER))}
chg = lambda g, k: G[g][1][k] / G[g][0][k] - 1
swing = {g: G[g][1]["contrib"] - G[g][0]["contrib"] for g in G}
net_swing = sum(swing.values())

# staggered rollout: group fare / control fare, indexed to Sep 24–Feb 25 (= 1.000)
def fare_index(routes, month):
    base = agg(routes, A)["fare"] / agg(OTHER, A)["fare"]
    m = (month, month)
    return agg(routes, m)["fare"] / agg(OTHER, m)["fare"] / base


fx = {"mar_before": fare_index(MAR, "2025-02"), "mar_after": fare_index(MAR, "2025-03"),
      "apr_before": fare_index(APR, "2025-03"), "apr_after": fare_index(APR, "2025-04")}

# each added trunk flight: tickets above the control trend vs tickets needed to cover its cost
tA, tB = G["trunk"]
oA, oB = G["other"]
added = tB["fl"] - tA["fl"]
extra_tix = (tB["tix"] - tA["tix"] * oB["tix"] / oA["tix"]) / added
inc_cost = tB["cost"] - tA["cost"] * (oB["cost"] / oB["fl"]) / (oA["cost"] / oA["fl"])  # $M at control cost growth
inc_rev = tB["rev"] - tA["rev"] * oB["rev"] / oA["rev"]
breakeven = inc_cost * 1e6 / added / tB["fare"]

# no maturation: Apr–Aug, identical trunk schedule in both years
mA, mB = agg(TRUNK, ("2025-04", "2025-08")), agg(TRUNK, ("2026-04", "2026-08"))
q = ",".join(f"'{r}'" for r in TRUNK)
slots = [set(db.execute(f"select route, substr(date, 6) from cost_events where route in ({q}) and month between ? and ?",
                        w).fetchall()) for w in (("2025-04", "2025-08"), ("2026-04", "2026-08"))]
assert mA["fl"] == mB["fl"] == len(slots[0] & slots[1])  # same route, month-day and departure time
apr_aug_tix = mB["tix"] / mA["tix"] - 1
new_fy = agg(NEW, ("2025-09", "2026-08"))
new_cov = new_fy["rev"] / new_fy["cost"]

# guard rails: the verified figures in HANDOFF.md / slides/OUTLINE.md "Verification notes"
assert abs(chg("trunk", "fl") - 1.008) < 0.002 and abs(chg("trunk", "rev") - 0.031) < 0.001
assert abs(chg("other", "rev") - 0.040) < 0.001 and abs(chg("trunk", "tpf") + 0.437) < 0.002
assert abs(tA["contrib"] - 1.15) < 0.01 and abs(tB["contrib"] + 23.76) < 0.01
assert abs(fx["mar_after"] - 0.909) < 0.001 and abs(fx["apr_before"] - 1.021) < 0.001
assert abs(fx["apr_after"] - 0.913) < 0.001 and abs(fx["mar_before"] - 1.007) < 0.001
assert abs(extra_tix - 1.9) < 0.05 and abs(breakeven - 19) < 0.5 and abs(apr_aug_tix) < 0.002
assert 0.31 < new_cov < 0.35


def r1(x):
    return Decimal(str(x)).quantize(Decimal("0.1"), ROUND_HALF_UP)


def pct(x, nd=1):
    """Signed percent, half-up, with a typographic minus: 0.0307 -> '+3.1%'."""
    q = Decimal(str(abs(x) * 100)).quantize(Decimal(1) if nd == 0 else Decimal("0." + "0" * (nd - 1) + "1"),
                                            ROUND_HALF_UP)
    return (MINUS if x < 0 else "+") + f"{q}%"


def usd(x, sign=True):
    """$M with sign outside the dollar: -23.76 -> '−$23.8M'."""
    m = money(x, sign)
    return (m[0] + "$" + m[1:] if m[0] in "+" + MINUS else "$" + m) + "M"


def drop_charts(s):
    """reset() removes shapes but not their chart parts; drop those relationships so reruns don't pile up."""
    for sh in list(s.shapes):
        if sh.has_chart:
            rid = sh._element.xpath(".//c:chart/@r:id")[0]
            sh._element.getparent().remove(sh._element)
            s.part.drop_rel(rid)


def no_frame(chart):
    """Transparent chart area with no border (sits on a white card)."""
    cs = chart._chartSpace
    sp = cs.find(qn("c:spPr"))
    if sp is None:
        sp = cs.makeelement(qn("c:spPr"), {})
        cs.find(qn("c:chart")).addnext(sp)  # schema order: chart, spPr, txPr, externalData
    for child in list(sp):
        sp.remove(child)
    sp.append(sp.makeelement(qn("a:noFill"), {}))
    ln = sp.makeelement(qn("a:ln"), {})
    ln.append(ln.makeelement(qn("a:noFill"), {}))
    sp.append(ln)


CATS = ["Doubled trunk", "New markets", "Other routes"]


def mini_chart(s, x, y, w, h, before, after, fmt, lo, hi):
    cd = CategoryChartData()
    cd.categories = CATS
    cd.add_series("Sep 2024–Feb 2025", before)
    cd.add_series("Sep 2025–Feb 2026", after)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    no_frame(ch)
    ch.has_legend = False
    ch.font.name, ch.font.size, ch.font.color.rgb = BODY, Pt(10), INK
    plot = ch.plots[0]
    plot.gap_width, plot.overlap, plot.vary_by_categories = 55, -6, False
    for i, ser in enumerate(plot.series):
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = BEFORE if i == 0 else NAVY
        dl = ser.data_labels
        dl.show_value = True
        dl.number_format, dl.number_format_is_linked = fmt, False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size, dl.font.bold = Pt(10), i == 1
        dl.font.color.rgb = MUTED if i == 0 else INK
    trunk_after = plot.series[1].points[0]
    trunk_after.format.fill.solid()
    trunk_after.format.fill.fore_color.rgb = ORANGE
    ca = ch.category_axis
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(10)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale, va.maximum_scale = lo, hi
    return ch


SO_WHAT_PT = 11
NB_ARROW = "\u00a0→\u00a0"  # keep "1.007 → 0.909" on one line
EM = {(HEAD, True): 0.63, (HEAD, False): 0.58, (BODY, True): 0.53, (BODY, False): 0.50}  # as slides/check_fit.py


def est_height(runs, width, size):
    """Wrapped height (in) of one paragraph of (text, overrides) runs; same glyph-width heuristic as check_fit."""
    line_w, lines = 0.0, 1
    for t, o in runs:
        em = EM[(o.get("font", BODY), o.get("bold", False))] * o.get("size", size) / 72
        for i, word in enumerate(t.split(" ")):
            w = (len(word) + (1 if i or line_w else 0)) * em
            if line_w + w > width and line_w > 0:
                lines, line_w = lines + 1, len(word) * em
            else:
                line_w += w
    return lines * size * 1.2 / 72


def scale(vals):
    """Value-axis limits that leave ~20% of the plot for data labels beyond the longest bar(s)."""
    top, bottom = max(vals), min(min(vals), 0)
    if bottom == 0:
        return 0, top / 0.8
    total = (top - bottom) / 0.6
    return bottom - 0.2 * total, top + 0.2 * total


def series(key, nd=1):
    vals = [(G[g][0][key], G[g][1][key]) for g in ("trunk", "new", "other")]
    rnd = (lambda v: int(round(v))) if nd == 0 else (lambda v: float(r1(v)))
    return [rnd(a) for a, _ in vals], [rnd(b) for _, b in vals]


TITLE = "Most of the loss came from doubling flights on existing trunk routes, not the 3 new markets"
SOURCE = ("Source: NorthByte transactions & costs, Sep–Feb FY24-25 vs FY25-26; revenue by purchase month, cost by "
          "flight date; before gates")


def build(s):
    drop_charts(s)
    reset(s, KICKER, TITLE, SOURCE, notes())
    t0, t1 = G["trunk"]
    n0, n1 = G["new"]
    o0, o1 = G["other"]

    # legend row
    text(s, 0.6, 1.68, 2.1, 0.26, [[("Sep–Feb, one year apart:", {"bold": True, "color": NAVY})]], size=11)
    box(s, 2.75, 1.75, 0.15, 0.15, fill=BEFORE)
    text(s, 2.96, 1.68, 2.2, 0.26, ["Sep 2024–Feb 2025 (before)"], size=10.5)
    box(s, 5.2, 1.75, 0.15, 0.15, fill=ORANGE)
    box(s, 5.38, 1.75, 0.15, 0.15, fill=NAVY)
    text(s, 5.59, 1.68, 2.96, 0.26, ["Sep 2025–Feb 2026 (trunk in orange)"], size=10.5)

    hi = lambda v: (v, {"bold": True, "color": ORANGE})
    panels = [
        ("Flights", series("fl", 0), '#,##0',
         [hi(f"trunk {pct(chg('trunk', 'fl'), 0)}"), (f"  ·  other routes {pct(chg('other', 'fl'))}", {})]),
        ("Revenue, $M", series("rev"), '0.0',
         [hi(f"trunk {pct(chg('trunk', 'rev'))}"), (f", less than other routes ({pct(chg('other', 'rev'))})", {})]),
        ("Tickets per flight", series("tpf"), '0.0',
         [hi(f"trunk {r1(t0['tpf'])} → {r1(t1['tpf'])} ({pct(chg('trunk', 'tpf'), 0)})"),
          (f"  ·  new markets {r1(n1['tpf'])}", {})]),
        ("Contribution before gates, $M", series("contrib"), f'"+"0.0;"{MINUS}"0.0;0.0',
         [hi(f"trunk {money(t0['contrib'])} → {money(t1['contrib'])}"),
          (f"  ·  new markets {money(n0['contrib'])} → {money(n1['contrib'])}", {})]),
    ]
    cw, chh, gx, gy, x0, y0 = 3.825, 1.95, 0.3, 0.3, 0.6, 2.02
    for k, (head, (before, after), fmt, sub) in enumerate(panels):
        cx, cy = x0 + (k % 2) * (cw + gx), y0 + (k // 2) * (chh + gy)
        box(s, cx, cy, cw, chh, fill=WHITE, line=LINE)
        text(s, cx + 0.14, cy + 0.08, cw - 0.28, 0.24, [[(head, {"bold": True, "color": NAVY, "font": HEAD})]],
             size=11)
        text(s, cx + 0.14, cy + 0.32, cw - 0.28, 0.22, [[(t.replace(" → ", NB_ARROW), o) for t, o in sub]], size=10.5)
        mini_chart(s, cx + 0.06, cy + 0.56, cw - 0.12, chh - 0.62, before, after, fmt, *scale(before + after))

    # key: name every market
    b = {"bold": True, "color": NAVY}
    text(s, 0.6, 6.33, 7.95, 0.55, [[
        ("Doubled trunk", b),
        (" (10 directions): Montréal–Toronto, Halifax–Toronto, Halifax–Montréal, Ottawa–Toronto, Montréal–Vancouver; "
         "each got an afternoon flight for every morning flight from Mar/Apr 2025.  ", {}),
        ("New markets", b), (" (6): London–Toronto (Feb 2025), Montréal–Ottawa, Moncton–Halifax.  ", {}),
        ("Other routes", b), (": the other 38 directions.", {})]], size=10.5)

    # SO WHAT panel
    px, py, pw, ph = 8.85, 1.68, 3.88, 5.2
    box(s, px, py, pw, ph, fill=PANEL)
    text(s, px + 0.25, py + 0.16, pw - 0.5, 0.26, [[("SO WHAT", {"bold": True, "color": ORANGE, "font": HEAD})]],
         size=11)
    items = [
        ("Frequency, not new markets.",
         f" Sep–Feb contribution fell ${r1(-net_swing)}M: doubled trunk {usd(swing['trunk'])}, new markets "
         f"{usd(swing['new'])}, other routes {usd(swing['other'])} (cost per flight "
         f"{pct((o1['cost'] / o1['fl']) / (o0['cost'] / o0['fl']) - 1, 0)})."),
        ("The added flights cut fares.",
         f" Each group's fare index vs other routes fell ~10% in the month its flights doubled: Montréal–Toronto "
         f"and Ottawa–Toronto in March ({fx['mar_before']:.3f} → {fx['mar_after']:.3f}); Halifax–Toronto, "
         f"Halifax–Montréal and Montréal–Vancouver in April ({fx['apr_before']:.3f} → {fx['apr_after']:.3f})."),
        ("Few new passengers.",
         f" Each added trunk flight carried ~{r1(extra_tix)} extra tickets; it needed ~{breakeven:.0f} to break even."),
        ("Waiting won't fix it.",
         f" Apr–Aug 2026 trunk tickets {pct(apr_aug_tix)} vs 2025, on an identical schedule. New markets still cover "
         f"only a third of their cost."),
    ]
    y, tw, size = py + 0.5, pw - 0.88, SO_WHAT_PT
    for i, (lead, body) in enumerate(items, start=1):
        runs = [(lead, {"bold": True, "color": NAVY}), (body.replace(" → ", NB_ARROW), {})]
        h = est_height(runs, tw, size) * 1.08
        label(box(s, px + 0.25, y + 0.02, 0.3, 0.3, fill=ORANGE, shape=MSO_SHAPE.OVAL), str(i), size=10)
        text(s, px + 0.68, y, tw, h, [runs], size=size)
        y += h + 0.1
    imp = [("→ Remove the added trunk flights first; the new markets are a smaller fix.",
            {"bold": True, "color": ORANGE})]
    ih = est_height(imp, pw - 0.45, size) * 1.08
    assert y + 0.05 < py + ph - 0.14 - ih, f"SO WHAT items run to {y:.2f}in"
    text(s, px + 0.25, py + ph - 0.14 - ih, pw - 0.45, ih, [imp], size=size, anchor=MSO_ANCHOR.BOTTOM)


def notes():
    t0, t1 = G["trunk"]
    n0, n1 = G["new"]
    o0, o1 = G["other"]
    t4 = [r for r in TRUNK if r not in ("YUL-YVR", "YVR-YUL")]
    four = (agg(t4, A), agg(t4, B))  # the 4 trunk markets without Montréal–Vancouver
    return f"""THE EXPANSION (guiding Q1/Q2). ~45 sec.
Say: "The expansion had two parts: three new markets, and a second rotation on five existing trunk markets. Compare the same six months a year apart. On the trunk, flights doubled ({pct(chg('trunk', 'fl'), 0)}) but revenue rose only {pct(chg('trunk', 'rev')).lstrip('+')}, less than on routes we did not touch ({pct(chg('other', 'rev'))}). Tickets per flight fell {pct(-chg('trunk', 'tpf'), 0).lstrip('+')}, and a {usd(t0['contrib'], False)} profit became a {usd(-t1['contrib'], False)} loss. The new markets lost {usd(-n1['contrib'], False)}: real, but about a fifth of the trunk swing."
Then the proof that the extra flights caused the fare cut: the rollout was staggered. Montréal–Toronto and Ottawa–Toronto doubled in March 2025 and their fare against other routes fell from {fx['mar_before']:.3f} to {fx['mar_after']:.3f} that month. Halifax–Toronto, Halifax–Montréal and Montréal–Vancouver were still at {fx['apr_before']:.3f} in March and fell to {fx['apr_after']:.3f} in April, when theirs doubled.
Close: "Each added flight brought about {r1(extra_tix)} extra passengers when it needed about {breakeven:.0f}, and a year later nothing has grown into the schedule. So the fix is to remove the added frequency, not to wait for it to fill."

Numbers (all recomputed by slides/build_expansion.py from analysis/airline_analysis.sqlite; Sep–Feb = 2024-09..2025-02 vs 2025-09..2026-02; revenue = all transaction types net of refunds by purchase month; flights and cost from cost_events by flight date; before gates):
- Doubled trunk, 10 directions (Ottawa–Toronto also got a midday flight, so "afternoon flight" is shorthand there): flights {t0['fl']:,} → {t1['fl']:,}; revenue ${r1(t0['rev'])}M → ${r1(t1['rev'])}M; tickets per flight {r1(t0['tpf'])} → {r1(t1['tpf'])}; contribution {usd(t0['contrib'])} → {usd(t1['contrib'])} ({usd(swing['trunk'])}). Average fare {pct(chg('trunk', 'fare'))}.
- New markets, 6 directions: flights {n0['fl']} → {n1['fl']:,} (the {n0['fl']} are London–Toronto's launch month, Feb 2025); revenue ${n0['rev']:.2f}M → ${r1(n1['rev'])}M; contribution {usd(n0['contrib'])} → {usd(n1['contrib'])}. FY25-26 cost coverage {new_cov:.1%} ("a third"; 32–34% in every month per HANDOFF).
- Other 38 directions: flights {o0['fl']:,} → {o1['fl']:,}; revenue {pct(chg('other', 'rev'))}; contribution {usd(o0['contrib'])} → {usd(o1['contrib'])} because cost per flight rose {pct((o1['cost'] / o1['fl']) / (o0['cost'] / o0['fl']) - 1)} (the Jan–Jul 2025 cost ramp, issue 2), not because of the schedule.
- Network Sep–Feb contribution fell ${r1(-net_swing)}M; the trunk is {swing['trunk'] / net_swing:.0%} of that, the new markets {swing['new'] / net_swing:.0%}.
- Fare index = group ticket revenue per ticket ÷ the same for the 38 other directions, divided by that ratio over Sep 24–Feb 25 (OUTLINE "Verification notes", additional data 1).
- Extra tickets per added flight: trunk tickets above the other routes' trend ({pct(oB['tix'] / oA['tix'] - 1)}), ÷ {added:,} added flights. Break-even: incremental cost ${r1(inc_cost)}M (trunk cost at the other routes' cost-per-flight growth) ÷ added flights ÷ FY25-26 Sep–Feb trunk fare (${tB['fare']:,.0f}). Incremental revenue vs the other routes' trend: {usd(inc_rev)} (OUTLINE additional data 2).
- Apr–Aug 2025 vs 2026: {mA['fl']:,} trunk flights in both years, identical route-date-time slots; tickets {mA['tix']:,} → {mB['tix']:,} ({pct(apr_aug_tix)}).
Q&A: slide 17 shows Montréal–Vancouver apart from the 4 other trunk markets because its fix differs (exit, not revert). It is in the doubled trunk here because it was doubled in April 2025 like the others. The 4 markets alone (8 directions): flights {pct(four[1]['fl'] / four[0]['fl'] - 1)}, revenue {pct(four[1]['rev'] / four[0]['rev'] - 1)}, tickets per flight {pct((four[1]['tix'] / four[1]['fl']) / (four[0]['tix'] / four[0]['fl']) - 1)}.
Concede if asked: the data cannot show which departure a passenger took, so AM and PM loads cannot be compared directly; the route-level test above is the evidence."""


def main():
    prs = Presentation(DECK)
    build(find_slide(prs, KICKER))
    prs.save(DECK)
    print(f"built {KICKER}: trunk swing {money(swing['trunk'])}M, new {money(swing['new'])}M, "
          f"other {money(swing['other'])}M; fare idx {fx}; extra tix {extra_tix:.2f} vs {breakeven:.1f}")


if __name__ == "__main__":
    main()
