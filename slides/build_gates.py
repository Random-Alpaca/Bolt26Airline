"""Fill the "4 · GATE LEASES" slide (guiding Q5): how gate lease costs changed from the 2024-25 to the 2025-26
lease year, what NorthByte should expect for 2026-27, gate cost per departing flight under the case rule, and the
loss after gates.

Run from the repo root:  uv run --no-project --with python-pptx python slides/build_gates.py
Edits the deck (slides/deckkit.py: DECK, or NB_DECK) in place and is safe to rerun: the slide is reset to its
frame (top bar, kicker, title, footer) before content is added. Every number is recomputed from
analysis/airline_analysis.sqlite and asserted against the verified figures (HANDOFF.md, OUTLINE.md
"Verification notes"; the teammate's outputs/tables/05_*.csv agree to the cent).

Method:
- Each physical gate (city, gate, lease year) is counted once: the file repeats a shared gate's cost on every
  route row, so the raw row sum ($17.4M for 2024-25) double-counts.
- Lease years run Jan 1 to Jan 1 (case): "2024-25" = calendar 2024, "2025-26" = 2025, "2026-27" = 2026 (lost).
- 2026-27 estimate: each gate repeats its own observed 2024-25 -> 2025-26 percentage change, same 34 gates.
  Range: every gate at the smallest or largest observed rise (+1.3% / +2.6%): $10.01M to $10.13M. Frozen rates
  ($9.88M) are only a floor: all 34 gates rose.
- Case rule: an airport's gate total is split evenly across all flights departing it (no arrival charge).
  FY25-26 gate cost is pro-rated by fiscal month (4 months of the 2025-26 lease + 8 of the 2026-27 estimate):
  $10.02M, giving the deck's -$58.4M after gates (analysis/alternatives.py uses the same basis). Per-departure
  figures use each airport's fiscal FY25-26 gate cost against its FY25-26 departures.
"""
import sqlite3
from decimal import Decimal

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from deckkit import (BODY, DECK, HEAD, INK, LINE, MINUS, MUTED, NAVY, ORANGE, PANEL, ROOT, WHITE, box, find_slide,
                     label, money, reset, text)

KICKER = "4 · GATE LEASES"
Y24, Y25 = "2024-2025", "2025-2026"
NAME = {"Montreal": "Montréal", "Quebec City": "Québec City"}
HUBS = ["Toronto", "Vancouver", "Montreal", "Calgary"]
MULTI = ["Halifax", "Ottawa", "Edmonton"]
EST_FILL = RGBColor.from_string("F4A97F")  # light orange: estimate, not actual
PRIOR_FILL = RGBColor.from_string("A9B6C9")

# ---------- data
db = sqlite3.connect(ROOT / "analysis/airline_analysis.sqlite")
gates = {}  # (city, gate) -> {year: cents, "assigned": bool}
for city, gate, year, hi, lo, n_routes in db.execute(
        "select city, gate, year, max(amount_cents), min(amount_cents), sum(route <> '') "
        "from gates group by city, gate, year"):
    assert hi == lo, (city, gate, year)  # repeated rows of one gate always agree
    g = gates.setdefault((city, gate), {})
    g[year] = hi
    g["assigned"] = g.get("assigned", False) or n_routes > 0
assert len(gates) == 34 and all(Y24 in g and Y25 in g for g in gates.values())
for g in gates.values():
    g["est"] = Decimal(g[Y25]) * g[Y25] / g[Y24]  # repeat this gate's own % change (cents)

row_sum24 = db.execute("select sum(amount_cents) from gates where year = ?", (Y24,)).fetchone()[0] / 100
L24 = sum(g[Y24] for g in gates.values()) / 100
L25 = sum(g[Y25] for g in gates.values()) / 100
E26 = float(sum(g["est"] for g in gates.values())) / 100
FLAT26, DOLLAR26 = L25, 2 * L25 - L24  # other scenarios: rates frozen; each gate repeats its $ change
FY2_FISCAL = L25 * 4 / 12 + E26 * 8 / 12  # Sep-Dec 2025 on the 2025-26 lease, Jan-Aug 2026 on the 2026-27 estimate
pct = [g[Y25] / g[Y24] - 1 for g in gates.values()]
LO26, HI26 = L25 * (1 + min(pct)), L25 * (1 + max(pct))  # every gate at the smallest / largest observed rise
IDLE = sorted(k for k, g in gates.items() if not g["assigned"])
IDLE24, IDLE25 = (sum(gates[g][y] for g in IDLE) / 100 for y in (Y24, Y25))
IDLE26 = float(sum(gates[g]["est"] for g in IDLE)) / 100

airports = {}
for (city, _), g in gates.items():
    a = airports.setdefault(city, {"n": 0, Y24: 0, Y25: 0, "est": Decimal(0), "idle": 0})
    a["n"] += 1
    a[Y24] += g[Y24] / 100
    a[Y25] += g[Y25] / 100
    a["est"] += g["est"] / 100
    a["idle"] += 0 if g["assigned"] else g[Y25] / 100
    a["fiscal"] = a.get("fiscal", 0) + (g[Y25] * 4 / 12 + float(g["est"]) * 8 / 12) / 100  # FY25-26 accrual
    a["idle_fiscal"] = a.get("idle_fiscal", 0) + (0 if g["assigned"] else (g[Y25] * 4 / 12 + float(g["est"]) * 8 / 12) / 100)
dep = dict(db.execute("select origin, count(*) from (select route, date, origin from costs where period = 2 "
                      "group by route, date) group by origin"))
DEPS = sum(dep.values())
DEPS1 = db.execute("select count(*) from (select 1 from costs where period = 1 group by route, date)").fetchone()[0]
assert set(dep) == set(airports) and DEPS == 15729 and DEPS1 == 13072
for c, a in airports.items():
    a["per_dep"] = a["fiscal"] / dep[c]
    a["idle_per_dep"] = a["idle_fiscal"] / dep[c]
AVG = FY2_FISCAL / DEPS

contrib = db.execute("select (select sum(amount_cents) from transactions where period = 2) - "
                     "(select sum(amount_cents) from costs where period = 2)").fetchone()[0] / 1e8  # $M
cost1, cost2 = (db.execute("select sum(amount_cents) from costs where period = ?", (p,)).fetchone()[0] / 1e8
                for p in (1, 2))
AFTER_LEASE = contrib - L25 / 1e6  # if 2026 rates were frozen at 2025 levels (floor)
AFTER = contrib - FY2_FISCAL / 1e6  # the deck's basis
HUB_SHARE_UP = sum(airports[c][Y25] - airports[c][Y24] for c in HUBS) / (L25 - L24)
HUB_LO, HUB_HI = (f(airports[c]["per_dep"] for c in HUBS) for f in (min, max))

# verified figures (HANDOFF.md, OUTLINE.md, key_issues_verification.md, teammate 05_*.csv)
assert (L24, L25, row_sum24) == (9_664_400, 9_877_700, 17_423_300)
assert abs(LO26 - 10_008_846.29) < 1 and abs(HI26 - 10_134_802.18) < 1
assert abs(E26 - 10_095_844.04) < 1 and abs(FY2_FISCAL - 10_023_129.36) < 1
assert [f"{c} {gt}" for c, gt in IDLE] == ["Halifax A3", "Ottawa A2"] and IDLE25 == 433_500
assert abs(contrib - -48.377) < 0.001 and money(AFTER) == f"{MINUS}58.4" and money(AFTER_LEASE) == f"{MINUS}58.3"
assert abs((cost2 - cost1) - 45.39) < 0.01
WORST = max(airports, key=lambda c: airports[c]["per_dep"])
assert WORST == "Whitehorse" and round(airports[WORST]["per_dep"]) == 1951
assert abs(sum(a["fiscal"] for a in airports.values()) - FY2_FISCAL) < 1


def k(x):  # whole dollars -> "$1,931"
    return f"${x:,.0f}"


def m2(x):  # dollars -> "$9.88M"
    return f"${x / 1e6:.2f}M"


# ---------- charts
def quiet(ch, size=10):
    ch.font.name, ch.font.size, ch.font.color.rgb = BODY, Pt(size), INK
    ch.has_title = False  # the panel header is the title; stops PowerPoint auto-titling
    ca, va = ch.category_axis, ch.value_axis
    ca.reverse_order = True
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.color.rgb = LINE
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(size)
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale = 0
    # show every category label (PowerPoint may otherwise skip some on a crowded axis)
    cat_ax = ch._chartSpace.find(".//" + qn("c:catAx"))
    if cat_ax.find(qn("c:tickLblSkip")) is None:
        skip = etree.SubElement(cat_ax, qn("c:tickLblSkip"))
        skip.set("val", "1")
        nm = cat_ax.find(qn("c:noMultiLvlLbl"))
        if nm is not None:
            nm.addprevious(skip)


def lease_chart(s, x, y, w, h):
    groups = [(NAME.get(c, c), [c]) for c in HUBS]
    groups.append(("Halifax, Ottawa, Edmonton", MULTI))
    rest = sorted(set(airports) - set(HUBS) - set(MULTI))
    assert len(rest) == 13 and all(airports[c]["n"] == 1 for c in rest)
    groups.append(("13 one-gate airports", rest))
    cd = CategoryChartData()
    cd.categories = [g for g, _ in groups]
    for key, nm in ((Y24, "2024-25 lease"), (Y25, "2025-26 lease"), ("est", "2026-27 est.")):
        cd.add_series(nm, [round(float(sum(airports[c][key] for c in cs)) / 1e6, 2) for _, cs in groups])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    quiet(ch)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10)
    plot = ch.plots[0]
    plot.gap_width = 30
    plot.overlap = 0
    ch.value_axis.maximum_scale = 2.75
    for ser, fill in zip(plot.series, (PRIOR_FILL, NAVY, EST_FILL)):
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = fill
        dl = ser.data_labels
        dl.show_value = True
        dl.number_format, dl.number_format_is_linked = "0.00", False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(10)
        dl.font.color.rgb = INK
    plot.series[2].data_labels.font.bold = True
    return groups


def per_dep_chart(s, x, y, w, h):
    rows = [(NAME.get(c, c), a["per_dep"], a["idle_per_dep"]) for c, a in airports.items()]
    rows.append((f"All {len(airports)} airports", AVG, 0.0))
    rows.sort(key=lambda r: -r[1])
    cd = CategoryChartData()
    cd.categories = [r[0] for r in rows]
    cd.add_series("Gate cost per departure", [round(r[1], 2) for r in rows])
    cd.add_series("of which unassigned gate", [round(r[2], 2) for r in rows])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    quiet(ch)
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = 40
    plot.overlap = 100  # idle-gate bar drawn over the base of the total bar
    ch.value_axis.maximum_scale = 2500
    total, idle = plot.series
    for ser, fill in ((total, NAVY), (idle, ORANGE)):
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = fill
    dl = total.data_labels
    dl.show_value = True
    dl.number_format, dl.number_format_is_linked = '"$"#,##0', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(10)
    dl.font.color.rgb = INK
    for i, (nm, v, iv) in enumerate(rows):
        pt = total.points[i]
        if nm.startswith("All "):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = MUTED
        if iv or nm.startswith("All ") or i == 0:
            lab = (f"{k(v)} ({k(iv)} unassigned)" if iv else f"{k(v)} average" if nm.startswith("All ") else k(v))
            tf = pt.data_label.text_frame
            tf.text = lab
            r = tf.paragraphs[0].runs[0]
            r.font.size, r.font.bold = Pt(10), True
            r.font.color.rgb = ORANGE if iv else INK
            pt.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    return rows


# ---------- slide
def drop_orphan_charts(s):
    """deckkit.reset() removes chart frames but not their relationships, so each rerun would leave the old chart
    parts in the file. Drop chart relationships that no element on this slide still references."""
    xml = s._element.xml
    for rId in [r.rId for r in s.part.rels.values() if r.reltype == RT.CHART]:
        if f'"{rId}"' not in xml:
            s.part.drop_rel(rId)


def card(s, x, y, w, h, num, head, body):
    box(s, x, y, w, h, fill=WHITE, line=LINE)
    label(box(s, x + 0.12, y + 0.1, 0.42, 0.42, fill=ORANGE, shape=MSO_SHAPE.OVAL), num, size=14)  # as slide 17
    text(s, x + 0.66, y + 0.09, w - 0.78, 0.3, [head], size=12, bold=True, color=NAVY, font=HEAD)
    text(s, x + 0.66, y + 0.4, w - 0.78, h - 0.46, [body], size=10.5)


def build(s):
    reset(s, KICKER,
          f"Gate leases rose only 2.2%, to {m2(L25)}; expect about ${E26 / 1e6:.1f}M for the 2026-27 lease year",
          f"Source: NorthByte gates, costs & transactions; team analysis. FY25-26 gates = {m2(FY2_FISCAL)}: "
          f"Sep–Dec 2025 lease + Jan–Aug 2026 estimate.",
          NOTES)
    drop_orphan_charts(s)
    top, bot = 1.72, 6.05
    h = bot - top
    w1, w2, gap = 3.6, 3.85, 0.3
    x1, x2 = 0.6, 0.6 + 3.6 + gap
    x3 = x2 + w2 + gap
    w3 = 12.73 - x3

    # 1 lease cost by airport
    box(s, x1, top, w1, h, fill=PANEL)
    text(s, x1 + 0.15, top + 0.1, w1 - 0.3, 0.62, [
        [("Gate lease cost by airport, $M", {"bold": True, "color": NAVY, "font": HEAD, "size": 12, "space_after": 2})],
        [(f"Total {m2(L24)} → {m2(L25)} → {m2(E26)} est.", {"bold": True, "size": 10.5})],
        [("34 gates, each counted once", {"size": 10, "color": MUTED})]], size=10.5)
    lease_chart(s, x1 + 0.05, top + 0.78, w1 - 0.1, h - 0.84)

    # 2 cost per departure (case rule)
    box(s, x2, top, w2, h, fill=PANEL)
    text(s, x2 + 0.15, top + 0.1, w2 - 0.3, 0.62, [
        [("Gate cost per departing flight", {"bold": True, "color": NAVY, "font": HEAD, "size": 12, "space_after": 2})],
        [("FY25-26: airport's gate cost ÷ its departures", {"size": 10, "color": MUTED})],
        [("Orange: ", {"size": 10, "bold": True, "color": ORANGE}),
         ("gates with no routes assigned", {"size": 10, "color": MUTED})]], size=10)
    per_dep_chart(s, x2 + 0.05, top + 0.72, w2 - 0.1, h - 0.76)

    # 3 so what
    w = airports["Whitehorse"]
    cards = [
        ("Leases +2.2%: not the problem",
         f"{m2(L24)} → {m2(L25)} (+${(L25 - L24) / 1e6:.2f}M) on the same 34 gates, each up "
         f"{min(pct) * 100:.1f}–{max(pct) * 100:.1f}%. Flight costs rose ${cost2 - cost1:.1f}M; gates are "
         f"~{L25 / (cost2 * 1e6 + L25):.0%} of all cost."),
        (f"2026-27: expect about ${E26 / 1e6:.1f}M",
         f"Method: each gate repeats its own 2025 % increase → {m2(E26)} (+{(E26 / L25 - 1) * 100:.1f}%). "
         f"Range {m2(LO26)}–{m2(HI26)}: every gate at the lowest or highest 2025 rise."),
        (f"Whitehorse: {k(w['per_dep'])} a departure",
         f"{w['per_dep'] / AVG:.0f}× the {k(AVG)} average (hubs {k(HUB_LO)}–{k(HUB_HI)}): one gate for "
         f"~{dep['Whitehorse'] / 52.18:.1f} departures a week. The fixed price of serving the North, not a reason "
         f"to cut."),
        (f"Loss after gates: {MINUS}${-AFTER:.1f}M",
         f"{MINUS}${-contrib:.1f}M less {m2(FY2_FISCAL)} of FY25-26 gates. Gates stay when flights go; "
         f"release the 2 unassigned gates (${IDLE25 / 1e6:.2f}M/yr) in Jan 2027."),
    ]
    cg = 0.12
    heights = [0.98, 0.98, 0.98, h - 3 * 0.98 - 3 * cg]
    y = top
    for i, ((head, body), ch) in enumerate(zip(cards, heights), start=1):
        card(s, x3, y, w3, ch, str(i), head, body)
        y += ch + cg

    # assumptions strip (the case asks teams to state them)
    box(s, 0.6, bot + 0.12, 12.13, 6.88 - bot - 0.12, fill=PANEL)
    text(s, 0.78, bot + 0.17, 11.8, 6.88 - bot - 0.22, [[
        ("Assumptions:  ", {"bold": True, "color": NAVY}),
        (f"each physical gate is counted once (the file repeats a shared gate on every route row: "
         f"${row_sum24 / 1e6:.1f}M row sum vs {m2(L24)})  ·  lease years run Jan 1–Jan 1: 2024-25 = calendar 2024, "
         f"2025-26 = 2025; 2026-27 (calendar 2026) records were lost, same 34 gates assumed  ·  case rule: each airport's gate total is split evenly across its "
         f"departing flights; arrivals pay nothing  ·  “idle” = no route assigned; station managers confirm before "
         f"a gate is released.", {})]], size=10.5, anchor=MSO_ANCHOR.MIDDLE)


NOTES = f"""GATE LEASES (guiding Q5). ~45 sec.
Say: "Leadership asked about gates. They are a small, steady cost: {m2(L25)}, up 2.2%. Expect about ${E26 / 1e6:.1f}M for the 2026-27 leases."
1. Change. Counting each physical gate once, leases went ${L24:,.0f} (2024-25 lease, calendar 2024) to ${L25:,.0f} (2025-26, calendar 2025): +${L25 - L24:,.0f}, +{(L25 / L24 - 1) * 100:.2f}%. Same 34 gates at 20 airports in both years; every gate rose {min(pct) * 100:.1f}% to {max(pct) * 100:.1f}%. The 4 hubs are {sum(airports[c][Y25] for c in HUBS) / L25:.0%} of the bill and {HUB_SHARE_UP:.0%} of the increase. Flight costs rose ${cost2 - cost1:.1f}M (FY24-25 to FY25-26); gates are about {L25 / (cost2 * 1e6 + L25):.0%} of FY25-26 recorded cost including gates. The raw file sums to ${row_sum24 / 1e6:.1f}M because a shared gate repeats on each route row.
2. 2026-27 estimate (the lease from Jan 1, 2026 to Jan 1, 2027; its records were lost). Each gate repeats its own 2025 percentage increase: ${E26:,.0f} (+{(E26 / L25 - 1) * 100:.1f}%). Range: every gate at the smallest observed rise (+{min(pct) * 100:.1f}%) gives ${LO26:,.0f}; every gate at the largest (+{max(pct) * 100:.1f}%) gives ${HI26:,.0f}. Repeating each gate's dollar increase gives ${DOLLAR26:,.0f}. Frozen rates (${FLAT26:,.0f}) are only a floor: all 34 gates rose in 2025. So: expect ~${E26 / 1e6:.1f}M, range {m2(LO26)} to {m2(HI26)}. This is a scenario range, not a confidence interval. (FY26-27 itself spans 4 months of the 2026-27 lease and 8 of the 2027-28 lease.)
3. Case rule: each airport's gate total is split evenly over every flight departing it (arrivals pay nothing). FY25-26 basis = each airport's fiscal gate cost (Sep–Dec 2025 lease + Jan–Aug 2026 estimate) ÷ its FY25-26 departures (same basis as the −$58.4M). Network ${AVG:,.0f} per departure ({DEPS:,} departures). Whitehorse ${airports['Whitehorse']['per_dep']:,.0f} ({dep['Whitehorse']} departures on one gate), Yellowknife ${airports['Yellowknife']['per_dep']:,.0f}, St. John's ${airports["St. John's"]['per_dep']:,.0f}; hubs ${HUB_LO:,.0f} (Toronto) to ${HUB_HI:,.0f} (Vancouver). Ottawa is lowest (${airports['Ottawa']['per_dep']:,.0f}) because Ottawa–Toronto was doubled and Ottawa–Montréal added ({dep['Ottawa']:,} departures). A year earlier it was ${L24 / DEPS1:,.0f} ({DEPS1:,} FY24-25 departures): spreading a fixed cost over 20% more flights lowered gate cost per flight, but did not make those flights pay.
   These are allocations of a fixed cost, not savings: the airports with the fewest departures per gate cost most per flight (Whitehorse {dep['Whitehorse']} and Yellowknife {dep['Yellowknife']} FY25-26 departures on one gate each), and they stay (all 20 airports are kept). Remoteness alone is not the driver: Abbotsford, Regina, Saskatoon, Winnipeg and Québec City also sit above Vancouver, while one-gate Kelowna, Victoria, London and Moncton are among the cheapest. If asked "why not drop Whitehorse's gate?": the gate is the price of serving that community, and only 2 of the 34 gates have no routes.
4. Unassigned gates: Halifax A3 and Ottawa A2 have no route assigned in either year: ${IDLE24:,.0f} (2024-25), ${IDLE25:,.0f} (2025-26), ≈${IDLE26:,.0f} at 2026-27 estimated rates. They add ${airports['Halifax']['idle_per_dep']:,.0f} to each Halifax departure and ${airports['Ottawa']['idle_per_dep']:,.0f} to each Ottawa departure. "Unassigned" does not prove "unused" (backup, irregular operations), so station managers confirm before release at the Jan 2027 renewal (recommendation, no-regret move).
5. Loss after gates: FY25-26 contribution −${-contrib:.2f}M minus {m2(FY2_FISCAL)} of FY25-26 gates (Sep–Dec 2025 on the 2025-26 lease, Jan–Aug 2026 on the 2026-27 estimate) = −${-AFTER:.2f}M, shown as −$58.4M across the deck. Freezing 2026 at 2025 rates would give −${-AFTER_LEASE:.2f}M; the $0.15M gap changes no conclusion.
   Why it matters: gates are fixed. Removing flights does not remove gate cost unless every route at a gate or airport goes. That is why issue sizes and the savings from cutting flights are stated before gates; each alternative's bottom line subtracts the same {m2(FY2_FISCAL)}, less any gates it releases. Right-Size's flight cuts release none; Shrink to Core's {MINUS}$3.4M already credits the gates it frees (it closes 6 of 20 airports); a full exit of the new markets would free $0.69M. Only releasing a gate saves money.
Sources: analysis/airline_analysis.sqlite (gates, costs, transactions); analysis/tables/gate_period.csv; HANDOFF.md "Data rules"; OUTLINE.md "Verification notes". Teammate's outputs/tables/05_forecast_scenarios.csv, 05_gate_change_forecast.csv, 05_fiscal_allocation_summary.csv and 05_unassigned_gate_leases.csv agree to the cent (their after-gates figure differs only because their revenue basis drops 759 out-of-window transactions)."""


def main():
    prs = Presentation(DECK)
    build(find_slide(prs, KICKER))
    prs.save(DECK)
    print(f"built Gate leases: {m2(L24)} -> {m2(L25)} -> est {m2(E26)} (fiscal FY25-26 {m2(FY2_FISCAL)}); "
          f"after gates {money(AFTER)}; avg ${AVG:,.0f}/departure")


if __name__ == "__main__":
    main()
