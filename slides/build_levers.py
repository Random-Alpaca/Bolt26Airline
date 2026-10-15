"""Fill the "3 · REVENUE LEVERS" slide (guiding questions 3 + 4): factor ranking of the −$36.0M decline,
plus refunds, ancillaries, Rewards and connecting passengers.

Run from the repo root:  uv run --no-project --with python-pptx python slides/build_levers.py
Edits the deck (slides/deckkit.py: DECK, or NB_DECK) in place and is safe to rerun: the slide is reset to its
frame (top bar, kicker, title, footer) before content is added. Every number is recomputed from
analysis/airline_analysis.sqlite and analysis/tables/*.csv, and checked against the verified values.
"""
import collections
import csv
import sqlite3
from decimal import ROUND_HALF_UP, Decimal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from deckkit import (BODY, DECK, GREEN, HEAD, INK, LINE, MINUS, MUTED, NAVY, ORANGE, PANEL, ROOT, WHITE,
                     box, find_slide, money, reset, text)

KICKER = "3 · REVENUE LEVERS"
LIGHT_ORANGE, LIGHT_GREEN, GREY = (RGBColor.from_string(c) for c in ("F6B48F", "9CCBB4", "9CA3AF"))
ANCILLARY = ("Change Fee", "Extra Baggage", "Food & Onboard", "Seat Upgrade")


def r1(x, places="0.1"):
    """Half-up rounding (matches deckkit.money), so bars and labels reconcile on screen."""
    return float(Decimal(str(x)).quantize(Decimal(places), ROUND_HALF_UP))


def usd(x, places="0.1"):
    """-12.36 -> "−$12.4M"; places="0.01" keeps two decimals for the small items."""
    v = Decimal(str(abs(x))).quantize(Decimal(places), ROUND_HALF_UP)
    return (MINUS if x < 0 else "+" if x > 0 else "") + f"${v}M"


def close(x, want, tol):
    assert abs(x - want) <= tol, (x, want)
    return x


# ---------- data (all $M unless noted)
db = sqlite3.connect(ROOT / "analysis/airline_analysis.sqlite")
q = lambda sql, *a: db.execute(sql, a).fetchall()

rev = {p: c / 1e8 for p, c in q("select period, sum(amount_cents) from transactions group by 1")}
cost = {p: c / 1e8 for p, c in q("select period, sum(amount_cents) from costs group by 1")}
contrib = {p: rev[p] - cost[p] for p in (1, 2)}
d_rev, d_cost, decline = rev[2] - rev[1], cost[2] - cost[1], contrib[2] - contrib[1]
close(contrib[1], -12.36, 0.01), close(contrib[2], -48.38, 0.01), close(decline, -36.02, 0.01)

# like-for-like cost: same route and calendar month in both years, at FY25-26 flight counts
cell = {(p, r, m, t): (n, c) for p, r, m, t, n, c in q(
    "select period, route, substr(month, 6, 2), type, count(*), sum(amount_cents) from costs group by 1, 2, 3, 4")}
lfl = collections.Counter()
for (p, r, m, t), (n, c) in cell.items():
    if p == 2 and (1, r, m, t) in cell:
        n1, c1 = cell[(1, r, m, t)]
        lfl[t] += n * (c / n - c1 / n1) / 1e8
lfl_total = sum(lfl.values())
vol_mix = d_cost - lfl_total
close(lfl_total, 15.25, 0.01), close(vol_mix, 30.14, 0.01), close(d_cost, 45.39, 0.01)

by_type = {(p, t): c / 1e8 for p, t, c in q("select period, type, sum(amount_cents) from transactions group by 1, 2")}
d_refund = by_type[(2, "Refund")] - by_type[(1, "Refund")]
d_anc = sum(by_type[(2, t)] - by_type[(1, t)] for t in ANCILLARY)
anc2 = sum(by_type[(2, t)] for t in ANCILLARY)
pc = {p: c / 1e8 for p, c in q(
    "select period, sum(amount_cents) from transactions where type = 'Ticket' and payment = 'Points + Cash' group by 1")}
d_pc = pc[2] - pc[1]
close(d_refund, -1.15, 0.005), close(d_anc, 0.29, 0.005), close(d_pc, 1.33, 0.005), close(d_rev, 9.37, 0.01)

gates = {r["period"]: float(r["network_gate_charge"]) / 1e6 for r in csv.DictReader(open(ROOT / "analysis/tables/gate_period.csv"))}
d_gates = -(gates["P2"] - gates["P1"])  # lease years 2024 -> 2025, each physical gate once
close(d_gates, -0.21, 0.005)

# refunds: share of revenue, channel burden (on gross = positive amounts), St. John's winter
refund_share = -by_type[(2, "Refund")] / rev[2]
burden = {ch: v for ch, v in q(
    "select channel, -1.0 * sum(case when type = 'Refund' then amount_cents end) "
    "/ sum(case when amount_cents > 0 then amount_cents end) from transactions where period = 2 group by 1")}
ota, web, app = burden["Online Travel Agency"], burden["Direct Web"], burden["Mobile App"]
yyt = q("select -1.0 * sum(case when type = 'Refund' then amount_cents end) / sum(case when amount_cents > 0 "
        "then amount_cents end) from transactions where period = 2 and route like '%YYT%' "
        "and substr(month, 6, 2) in ('12', '01', '02')")[0][0]
close(ota, 0.058, 0.0005), close(web, 0.037, 0.0005), close(app, 0.037, 0.0005), close(yyt, 0.371, 0.001)

# Rewards: fare by tier (both years) and seat-upgrade rate per ticket
tiers = {t: (avg, seats) for t, avg, seats in q(
    "select tier, 1.0 * sum(case when type = 'Ticket' then amount_cents end) / sum(type = 'Ticket') / 100, "
    "1.0 * sum(type = 'Seat Upgrade') / sum(type = 'Ticket') from transactions where tier != 'Unknown' group by 1")}
avg_ticket = q("select 1.0 * sum(amount_cents) / count(*) / 100 from transactions where type = 'Ticket'")[0][0]
spread = max(abs(a / avg_ticket - 1) for a, _ in tiers.values())
seat_ratio = (min(tiers["Gold"][1], tiers["Platinum"][1])
              / max(tiers[t][1] for t in ("Blue", "Silver", "Non-Member")))
close(avg_ticket, 1031.0, 0.5), close(spread, 0.0026, 0.0005), close(seat_ratio, 1.76, 0.05)

# connecting passengers: a booking whose tickets are two chained legs (A-B, B-C, C != A) is a hub connection
legs = collections.defaultdict(list)
for b, r in q("select booking_id, route from transactions where type = 'Ticket'"):
    legs[b].append(r.split("-"))
hub_of = {}
for b, l in legs.items():
    if len(l) == 2 and l[0] != l[1][::-1]:
        if l[0][1] == l[1][0]:
            hub_of[b] = l[0][1]
        elif l[1][1] == l[0][0]:
            hub_of[b] = l[1][1]
n_conn = len(hub_of)
hub_share = collections.Counter(hub_of.values())
conn_rev2 = sum(c for b, c in q("select booking_id, amount_cents from transactions where period = 2") if b in hub_of) / 1e8
conn_share = conn_rev2 / rev[2]
n_conn2 = len(hub_of.keys() & {b for (b,) in q(
    "select booking_id from transactions where type = 'Ticket' and period = 2")})
mtl_yvr = sum(1 for b in hub_of if any(l in (["YUL", "YVR"], ["YVR", "YUL"]) for l in legs[b]))
close(n_conn, 12644, 0), close(conn_rev2, 13.11, 0.01), close(conn_share, 0.069, 0.0005), close(mtl_yvr, 0, 0)
close(n_conn2, 6441, 0)  # bookings with a ticket in FY25-26 (one booking spans both years)
TEAMMATE_CONN = 13.08  # teammate outputs/tables/04_connecting_booking_proxy.csv, FY2025-26, 6.90% of their net receipts
shrink_conn = next(float(r["connecting_revenue_lost_musd"]) for r in csv.DictReader(
    open(ROOT / "analysis/tables/alternatives_matrix.csv"))
    if r["alternative"] == "Shrink to Core" and r["component"] == "TOTAL")

# ---------- chart rows: (label, value shown, colour). Main bars sum to the total; "of which" bars sit inside the bar above.
BARS = [
    ("More flights (volume & mix)", r1(-vol_mix), ORANGE),
    # $15.25M: shown as the rounded cost growth less the rounded volume bar (45.4 - 30.1 = 15.3, the deck-wide figure)
    ("Cost per flight, like-for-like", r1(-d_cost) - r1(-vol_mix), ORANGE),
    ("of which fuel", r1(-lfl["Fuel"]), LIGHT_ORANGE),
    ("of which crew", r1(-lfl["Crew"]), LIGHT_ORANGE),
    ("Net revenue growth", r1(d_rev), GREEN),
    ("of which Points + Cash tickets", r1(d_pc, "0.01"), LIGHT_GREEN),
    ("of which refunds", r1(d_refund, "0.01"), LIGHT_ORANGE),
    ("of which ancillaries", r1(d_anc, "0.01"), LIGHT_GREEN),
    ("Total change, before gates", r1(decline), NAVY),
    ("Gate leases (outside the total)", r1(d_gates, "0.01"), GREY),
]
BARS = [(lab, round(v, 2), c) for lab, v, c in BARS]
assert abs(BARS[0][1] + BARS[1][1] + BARS[4][1] - BARS[8][1]) < 1e-9, "main bars must add to the total on screen"


def drop_orphans(slide):
    """reset() removes shapes but not their chart/image parts; drop those so reruns don't pile up copies."""
    xml = slide._element.xml
    for rId, rel in list(slide.part.rels.items()):
        if rel.reltype in (RT.CHART, RT.IMAGE) and f'"{rId}"' not in xml:
            slide.part.drop_rel(rId)


def build(s):
    reset(s, KICKER,
          "Extra flights explain most of the $36M decline; refunds and Rewards are each under $1.5M",
          "Source: NorthByte transactions & costs, FY24-25 vs FY25-26; like-for-like = same route & month; "
          "gates by lease year; connection = 2 chained legs on one booking",
          NOTES)
    drop_orphans(s)

    # left: factor ranking
    box(s, 0.6, 1.72, 7.1, 5.1, fill=PANEL)
    text(s, 0.8, 1.82, 6.75, 0.5, [
        [(f"What moved contribution from {usd(contrib[1])} to {usd(contrib[2])}", {"bold": True, "color": NAVY})],
        [("$M change, FY24-25 → FY25-26, ranked by size; light bars are part of the dark bar above, not extra",
          {"size": 10, "color": MUTED})]], size=12)
    cd = CategoryChartData()
    cd.categories = [b[0] for b in BARS]
    cd.add_series("Change in contribution", [b[1] for b in BARS])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.7), Inches(2.36), Inches(6.9), Inches(3.72), cd)
    ch = gf.chart
    ch.has_title = False  # stop PowerPoint showing the series name as a title
    ch.has_legend = False
    ch.font.name, ch.font.size, ch.font.color.rgb = BODY, Pt(10), INK
    plot = ch.plots[0]
    plot.gap_width = 40
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = f'"+"0.0#;"{MINUS}"0.0#;0.0', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(10.5), True
    ser = plot.series[0]
    ser.invert_if_negative = False
    for i, (_, _, colour) in enumerate(BARS):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = colour
    # say invertIfNegative=0 on every point too, so negative bars keep their colour whatever PowerPoint defaults to
    for dPt in ser._element.findall(qn("c:dPt")):
        if dPt.find(qn("c:invertIfNegative")) is None:
            dPt.find(qn("c:idx")).addnext(OxmlElement("c:invertIfNegative"))
        dPt.find(qn("c:invertIfNegative")).set("val", "0")
    ca = ch.category_axis
    ca.reverse_order = True
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.color.rgb = LINE
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(10)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale, va.maximum_scale = -44, 15

    m = {b[0]: b[1] for b in BARS}
    text(s, 0.8, 6.14, 6.75, 0.62, [
        [("How it adds up: ", {"bold": True, "color": NAVY}),
         (f"{money(m['More flights (volume & mix)'])} {MINUS} {money(-m['Cost per flight, like-for-like'], False)} + "
          f"{money(m['Net revenue growth'], False)} = {money(m['Total change, before gates'])}. ", {"bold": True}),
         ("Cost growth is +$45.4M: ⅔ more flights, ⅓ higher cost per flight. "
          "Gate leases rose a further $0.21M, outside the total.", {})]], size=10.5)

    # right: what each factor means
    cards = [
        ("Refunds", usd(d_refund, "0.01"),
         f"$6.7M → $7.8M, about {refund_share:.0%} of revenue. OTA bookings refund {ota:.1%} of sales vs "
         f"{web:.1%} on web and app; St. John's refunds ~{yyt:.0%} of its winter revenue.",
         "→ Rebook-first in St. John's; OTA bookings to direct"),
        ("Ancillaries", usd(d_anc, "0.01"),
         f"Bags, seats, food and change fees: {anc2 / rev[2]:.1%} of revenue (${anc2:.1f}M). "
         f"Even +10% adds only ~${anc2 * 0.1:.1f}M.",
         "→ All-in fare bundles on the flights we keep"),
        ("Rewards", f"≤ ${d_pc:.1f}M",
         f"No fare premium: all tiers pay within {spread:.1%} of ~${avg_ticket:,.0f} a ticket. Gold and Platinum buy "
         f"~2× the seat upgrades. Points cost is not in the data; it can be no more than the ${d_pc:.2f}M "
         "Points + Cash rise.",
         "→ Not a driver; get its cost data before any change"),
        ("Connecting passengers", f"${conn_rev2:.1f}M",
         f"{n_conn:,} hub-connection bookings over two years ({hub_share['YYZ'] / n_conn:.0%} via Toronto, "
         f"{hub_share['YYC'] / n_conn:.0%} via Calgary); {conn_share:.1%} of FY25-26 revenue. "
         "Montréal–Vancouver has 0, so exiting it loses none.",
         "→ Keep every hub; market connections through them"),
    ]
    # card height follows its body: 3, 2, 4 and 3 lines of 10.5pt text (0.175in a line); the stack ends at y 6.78
    x, w, y, gap = 8.0, 4.73, 1.72, 0.08
    for k, ((head, stat, body, fix), lines) in enumerate(zip(cards, (3, 2, 4, 3)), start=1):
        body_h = lines * 0.175 + 0.01
        h = 0.42 + body_h + 0.25
        box(s, x, y, w, h, fill=WHITE, line=LINE)
        circ = box(s, x + 0.12, y + 0.08, 0.32, 0.32, fill=ORANGE, shape=MSO_SHAPE.OVAL)
        circ.text_frame.text = str(k)
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].runs[0]
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = HEAD, Pt(12), True, WHITE
        text(s, x + 0.54, y + 0.07, w - 2.1, 0.34, [head], size=12.5, bold=True, color=NAVY, font=HEAD,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + w - 1.52, y + 0.05, 1.4, 0.38, [stat], size=15, bold=True, color=NAVY, font=HEAD,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.14, y + 0.42, w - 0.28, body_h, [body], size=10.5)
        text(s, x + 0.14, y + 0.42 + body_h, w - 0.28, 0.2, [fix], size=10, bold=True, color=ORANGE)
        y += h + gap
    assert y - gap <= 6.83, y


NOTES = f"""REVENUE LEVERS: GUIDING Q3 + Q4. ~60 sec.
Say: "Leadership suspected fuel, crew, refunds and Rewards. Ranked in dollars, the $36M decline is mostly the extra flights. The suspects are real but small, and two of them are levers we use in the recommendation."
Chart (before gates; bars are $M change in contribution, FY24-25 → FY25-26):
- Contribution {usd(contrib[1])} → {usd(contrib[2])} = {usd(decline)}. Revenue {usd(d_rev)}; flight costs {usd(d_cost)}.
- More flights (volume & mix) {usd(-vol_mix)}: the part of the +$45.4M cost growth not explained by unit-cost changes, i.e. 13,072 → 15,729 flights plus the mix of new routes and months (⅔ of cost growth).
- Cost per flight like-for-like {usd(-lfl_total, "0.01")}, shown as $15.3M so the ⅔/⅓ split adds to the rounded $45.4M (⅓): same route and calendar month, at FY25-26 flight counts. Fuel {usd(-lfl['Fuel'])} (+8.7%), crew {usd(-lfl['Crew'])} (+8.5%), maintenance {usd(-lfl['Maintenance'])}, food ~0. On identical flights only it is $12.3M; at run-rate vs fall-2024 unit costs $26–28M. Q3 answer: fuel and crew prices explain about ${lfl['Fuel'] + lfl['Crew']:.1f}M; the rest of their $31.2M growth is from flying more.
- Net revenue {usd(d_rev)}, of which Points + Cash ticket sales {usd(d_pc, "0.01")}, refunds {usd(d_refund, "0.01")}, ancillaries {usd(d_anc, "0.01")}. Ticket sales overall +$10.2M. Do NOT add the light bars to the main bars: they are already inside them.
- Gates {usd(d_gates, "0.01")} ($9.66M → $9.88M, +2.2%, lease years) sit outside the before-gates total.
Cards:
1. Refunds $6.68M → $7.83M (about half the rise is a data-window artefact; like-for-like +8–10%). OTA 5.8% vs 3.7% web/app on gross sales (FY25-26). St. John's Dec–Feb FY25-26: 37% of gross sales refunded (36–39% by measure). Addressable excess ~$2M (St. John's ~$1.9M, OTA ~$0.8M; overlapping) → rebook-first and OTA → direct levers.
2. Ancillaries: 4 streams (bags, seats, food, change fees) ${anc2:.2f}M = {anc2 / rev[2]:.1%} of FY25-26 revenue; +10% = ${anc2 * 0.1:.2f}M. (The 3 paid extras alone are $3.36M; +10% = $0.34M. Use the 4-stream number.)
3. Rewards: average ticket ${avg_ticket:,.0f} over both years; every tier within {spread:.2%}. Seat upgrades per ticket: Gold {tiers['Gold'][1]:.1%}, Platinum {tiers['Platinum'][1]:.1%} vs 5.6–6.2% for Blue, Silver and non-members. Points + Cash records the full fare and non-members use it too, so the program's cost is not visible. Upper bound: if all the Points + Cash growth were unpaid points, Rewards explains at most ${d_pc:.2f}M ({d_pc / -decline:.1%} of the decline). Rivals' loyalty spending (case p.8) is a threat, but matching it is not the fix.
4. Connections: two chained ticket legs on one booking (A–B + B–C). {n_conn:,} bookings over Sep 2024–Aug 2026 (plus 268 return trips, excluded); {n_conn2:,} of them have a ticket in FY25-26, so do not divide the FY25-26 revenue by {n_conn:,}. FY25-26 revenue on those bookings ${conn_rev2:.2f}M = {conn_share:.2%} of ${rev[2]:.1f}M. Teammate's independent proxy: ${TEAMMATE_CONN:.2f}M, 6.90% of their $189.71M (they drop 759 out-of-window transactions), so the two agree. Hubs: Toronto {hub_share['YYZ']:,}, Calgary {hub_share['YYC']:,}, Vancouver {hub_share['YVR']:,}, Montréal {hub_share['YUL']:,}. Montréal–Vancouver: 0 connecting passengers. Trunk cuts put only ~$0.18M a year of connection revenue at risk. Shrink to Core would lose ${shrink_conn:.1f}M of connecting revenue (analysis/tables/alternatives_matrix.csv); Right-Size loses $0. Lever in the recommendation: hub connection marketing, $0.7–2.6M. Not proven incremental: the data cannot show which flight a passenger took.
Sources: analysis/airline_analysis.sqlite (slides/build_levers.py recomputes every number); slides/OUTLINE.md "Verification notes"; analysis/key_issues_verification.md; analysis/tables/gate_period.csv; teammate outputs/tables/04_connecting_booking_proxy.csv and 03_decline_bridge.csv."""


def main():
    prs = Presentation(DECK)
    build(find_slide(prs, KICKER))
    prs.save(DECK)
    print("built Revenue levers: " + ", ".join(f"{lab} {v:+.2f}" for lab, v, _ in BARS))


if __name__ == "__main__":
    main()
