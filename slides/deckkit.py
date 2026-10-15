"""Shared drawing helpers for the slide builders (slides/build_*.py).

Every builder edits the deck in place. Set NB_DECK=/path/copy.pptx to build into a scratch copy instead.
Template frame on content slides: top bar, kicker (TextBox 2, Montserrat 11 bold orange), title (TextBox 3,
Montserrat 24 bold navy, 0.60/0.62, 12.10 wide, up to 2 lines), footer rule at y=7.00, source (TextBox 5,
Lato 9 muted), page number (TextBox 6). Content area: x 0.60–12.73 in, y ~1.65–6.90 in. Body text is Lato.
"""
import os
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK = Path(os.environ.get("NB_DECK", ROOT / "slides" / "NorthByte_FirstByte2026_template.pptx"))
NAVY, INK, MUTED, ORANGE, PANEL, LINE, WHITE, GREEN = (RGBColor.from_string(c) for c in (
    "0B1F3A", "1F2937", "6B7280", "E8692A", "F3F5F8", "D1D5DB", "FFFFFF", "2E7D5B"))
HEAD, BODY = "Montserrat", "Lato"
FRAME = {"Rectangle 1", "TextBox 2", "TextBox 3", "Rectangle 4", "TextBox 5", "TextBox 6"}
MINUS = "−"


def money(x, sign=True):
    # half-up, so -58.25 shows as -58.3 like the rest of the deck (f-strings round half-to-even)
    s = str(Decimal(str(abs(x))).quantize(Decimal("0.1"), ROUND_HALF_UP))
    return (MINUS if x < 0 else "+" if sign and x > 0 else "") + s


def find_slide(prs, kicker):
    """The slide whose kicker (TextBox 2) text equals `kicker`; robust to slides being added or deleted."""
    for s in prs.slides:
        if any(sh.name == "TextBox 2" and sh.has_text_frame and sh.text_frame.text == kicker for sh in s.shapes):
            return s
    raise KeyError(kicker)


def text(slide, x, y, w, h, paras, size=12, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=0.0, name=None):
    """paras: list of paragraphs; each is a str or a list of (text, {overrides}) runs.
    Run overrides: font, size, bold, italic, color; space_after (pt) on a paragraph's first run."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name, f.size, f.bold = o.get("font", font), Pt(o.get("size", size)), o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
        if isinstance(para, list) and para and "space_after" in para[0][1]:
            p.space_after = Pt(para[0][1]["space_after"])
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s


def label(shape, value, size=10, color=WHITE, bold=True, font=HEAD):
    """Centered text inside a shape (badges, pills, numbered circles)."""
    tf = shape.text_frame
    tf.text = value
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = font, Pt(size), bold, color
    return shape


def rule(slide, x1, y1, x2, y2, color=LINE, width=1.0, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    return c


def set_runs(shape, value):
    """Replace a text box's text, keeping the first run's formatting; drops other paragraphs/runs."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    p.runs[0].text = value
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def reset(slide, kicker, title, source, notes):
    """Strip a content slide back to its frame, then set kicker, title, source line and speaker notes."""
    for sh in list(slide.shapes):
        if sh.name not in FRAME:
            sh._element.getparent().remove(sh._element)
    for name, value in (("TextBox 2", kicker), ("TextBox 3", title), ("TextBox 5", source)):
        set_runs(next(s for s in slide.shapes if s.name == name), value)
    slide.notes_slide.notes_text_frame.text = notes
