"""Last step after the slide builders: drop the empty ROOT CAUSE slide (folded into Key issues, see OUTLINE.md)
and renumber the page boxes so they match slide order. Safe to rerun.

uv run --no-project --with python-pptx python slides/finalize.py
"""
from pptx import Presentation

from deckkit import DECK, find_slide, set_runs

prs = Presentation(DECK)
try:
    s = find_slide(prs, "ROOT CAUSE")
    ids = prs.slides._sldIdLst
    sid = next(x for x in ids if prs.slides.get(int(x.get("id"))) is s)
    prs.part.drop_rel(sid.rId)
    ids.remove(sid)
    print("removed ROOT CAUSE slide")
except KeyError:
    pass
for n, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.name == "TextBox 6" and sh.has_text_frame and sh.text_frame.text.strip().isdigit():
            set_runs(sh, str(n))
prs.save(DECK)
print(f"{len(prs.slides)} slides, page numbers updated")
